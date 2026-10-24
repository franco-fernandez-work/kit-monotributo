import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "private", "vendedores.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
AZUL = "1F4E79"; CLARO = "DDEBF7"; INPUT = "FFF2CC"
H = Font(bold=True, color="FFFFFF"); B = Font(bold=True)
fill = lambda c: PatternFill("solid", fgColor=c)
thin = Side(style="thin", color="BFBFBF"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
ARS = '"$" #,##0'; PCT = "0.0%"

wb = Workbook()

def title(ws, text, sub):
    ws["A1"] = text; ws["A1"].font = Font(bold=True, size=16, color=AZUL)
    ws["A2"] = sub; ws["A2"].font = Font(italic=True, color="595959")

ws = wb.active; ws.title = "Inicio"
title(ws, "Calculadora de Precios para Vendedores 2026", "Mercado Libre, tienda propia y venta directa. Funciona en Excel y Google Sheets.")
ws.column_dimensions["A"].width = 115
lines = [
 "CÓMO USARLO",
 "1. Pestaña 'Costos': cargá una vez tus comisiones e impuestos (celdas amarillas). Mirá los valores en tu cuenta de Mercado Libre,",
 "   en 'Facturación' o en el detalle de cada venta, porque cambian según categoría y período.",
 "2. Pestaña 'Productos': cargá cada producto con su costo, envío a tu cargo y el margen que querés ganar.",
 "   La planilla te dice a qué precio publicarlo en Clásica, Premium y en venta directa, y cuánto ganás de verdad.",
 "3. Si ya tenés un precio publicado, cargalo en 'Precio actual' y ves tu ganancia neta y margen real.",
 "4. Pestaña 'Simulador': probá rápido un producto y compará los tres canales lado a lado.",
 "",
 "Celdas amarillas = las cargás vos. Todo lo demás se calcula solo.",
 "",
 "CÓMO CALCULA",
 "Precio = (costo + envío + costo fijo) / (1 - comisión - cuotas - impuestos - margen deseado).",
 "Si el precio queda por debajo del umbral de costo fijo, se suma el costo fijo por unidad que cargues.",
 "",
 "AVISO",
 "Herramienta de cálculo personal. Las comisiones y retenciones las define cada plataforma y pueden cambiar: verificalas en tu cuenta.",
 "Para uso personal. Prohibida su reventa o redistribución.",
]
for i, l in enumerate(lines, 4):
    ws.cell(row=i, column=1, value=l)
    if l in ("CÓMO USARLO", "CÓMO CALCULA", "AVISO"): ws.cell(row=i, column=1).font = Font(bold=True, color=AZUL, size=12)

c = wb.create_sheet("Costos")
title(c, "Comisiones e impuestos", "Valores de ejemplo: reemplazalos por los de tu cuenta y categoría.")
c.column_dimensions["A"].width = 52; c.column_dimensions["B"].width = 16; c.column_dimensions["C"].width = 60
params = [
 ("ComClasica", "Comisión Mercado Libre publicación Clásica", 0.13, PCT, "Depende de la categoría. Mirala al publicar."),
 ("ComPremium", "Comisión Mercado Libre publicación Premium", 0.18, PCT, "Incluye cuotas sin interés según la categoría."),
 ("Umbral", "Precio por debajo del cual se cobra costo fijo", 33000, ARS, "Mercado Libre cobra un cargo fijo a productos baratos."),
 ("CostoFijo", "Costo fijo por unidad (productos baratos)", 1100, ARS, "Usá el valor de tu tramo de precio."),
 ("ComDirecta", "Comisión Mercado Pago en venta directa / tienda propia", 0.0629, PCT, "Depende del plazo de acreditación que elijas."),
 ("Cuotas", "Costo extra por ofrecer cuotas sin interés (Clásica)", 0.0, PCT, "Dejá 0 si no ofrecés cuotas en Clásica."),
 ("IIBB", "Ingresos Brutos / retenciones provinciales", 0.03, PCT, "Según tu provincia y régimen."),
 ("OtrasRet", "Otras retenciones o impuestos sobre la venta", 0.0, PCT, "Por ejemplo, impuesto a débitos y créditos."),
]
for r, (name, lab, val, fmt, note) in enumerate(params, 4):
    c.cell(row=r, column=1, value=lab).font = B
    cell = c.cell(row=r, column=2, value=val); cell.number_format = fmt; cell.fill = fill(INPUT); cell.border = box
    c.cell(row=r, column=3, value=note).font = Font(color="595959")
    wb.defined_names[name] = DefinedName(name, attr_text=f"'Costos'!$B${r}")

def precio(base, com, cell_margen, fijo=True):
    k = f"({com}+IIBB+OtrasRet+{cell_margen})"
    p1 = f"({base})/(1-{k})"
    p2 = f"(({base})+CostoFijo)/(1-{k})"
    if not fijo:
        return f'=IF({k}>=1,"Margen imposible",{p1})'
    return f'=IF({k}>=1,"Margen imposible",IF({p1}<Umbral,{p2},{p1}))'

def neto(precio_cell, base, com):
    return f'=IF(ISNUMBER({precio_cell}),{precio_cell}*(1-{com}-IIBB-OtrasRet)-({base})-IF({precio_cell}<Umbral,CostoFijo,0),"")'

p = wb.create_sheet("Productos")
title(p, "Precios por producto", "Cargá costo, envío y margen deseado. El resto se calcula solo.")
heads = ["Producto", "Costo unitario", "Envío a tu cargo", "Margen deseado", "Precio Clásica", "Precio Premium", "Precio venta directa",
         "Precio actual (opcional)", "Canal del precio actual", "Ganancia neta", "Margen real"]
widths = [30, 15, 15, 14, 16, 16, 18, 18, 18, 15, 13]
for i, (t, w) in enumerate(zip(heads, widths), 1):
    cell = p.cell(row=4, column=i, value=t); cell.font = H; cell.fill = fill(AZUL); cell.border = box
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    p.column_dimensions[chr(64 + i)].width = w
p.row_dimensions[4].height = 32
dv = DataValidation(type="list", formula1='"Clásica,Premium,Directa"', allow_blank=True); p.add_data_validation(dv)
N = 500
for r in range(5, 5 + N):
    base = f"B{r}+C{r}"
    for col in "ABCDH": p[f"{col}{r}"].fill = fill(INPUT)
    p[f"I{r}"].fill = fill(INPUT); dv.add(f"I{r}")
    blank = f'IF(B{r}="","",'
    p[f"E{r}"] = precio(base, "ComClasica+Cuotas", f"D{r}").replace("=IF(", f"={blank}IF(", 1) + ")"
    p[f"F{r}"] = precio(base, "ComPremium", f"D{r}").replace("=IF(", f"={blank}IF(", 1) + ")"
    p[f"G{r}"] = precio(base, "ComDirecta", f"D{r}", False).replace("=IF(", f"={blank}IF(", 1) + ")"
    com = f'IF(I{r}="Premium",ComPremium,IF(I{r}="Directa",ComDirecta,ComClasica+Cuotas))'
    p[f"J{r}"] = f'=IF(OR(H{r}="",B{r}=""),"",H{r}*(1-{com}-IIBB-OtrasRet)-({base})-IF(AND(I{r}<>"Directa",H{r}<Umbral),CostoFijo,0))'
    p[f"K{r}"] = f'=IF(J{r}="","",J{r}/H{r})'
    for col, fmt in zip("BCDEFGHJK", [ARS, ARS, PCT, ARS, ARS, ARS, ARS, ARS, PCT]):
        p[f"{col}{r}"].number_format = fmt
    for col in "ABCDEFGHIJK": p[f"{col}{r}"].border = box
p["A5"] = "Ejemplo: auriculares"; p["B5"] = 8000; p["C5"] = 0; p["D5"] = 0.25; p["H5"] = 15999; p["I5"] = "Clásica"
p.freeze_panes = "B5"
p.conditional_formatting.add(f"J5:J{4+N}", CellIsRule(operator="lessThan", formula=["0"], fill=fill("F8CBAD")))
p.conditional_formatting.add(f"K5:K{4+N}", CellIsRule(operator="lessThan", formula=["0.1"], fill=fill("F8CBAD")))

s = wb.create_sheet("Simulador")
title(s, "Simulador rápido", "Probá un producto y compará los tres canales.")
s.column_dimensions["A"].width = 34
for col in "BCD": s.column_dimensions[col].width = 18
for lab, r, val, fmt in (("Costo unitario", 4, 10000, ARS), ("Envío a tu cargo", 5, 0, ARS), ("Margen deseado", 6, 0.3, PCT)):
    s[f"A{r}"] = lab; s[f"A{r}"].font = B
    s[f"B{r}"] = val; s[f"B{r}"].number_format = fmt; s[f"B{r}"].fill = fill(INPUT); s[f"B{r}"].border = box
for col, t in zip("BCD", ["Clásica", "Premium", "Venta directa"]):
    s[f"{col}8"] = t; s[f"{col}8"].font = H; s[f"{col}8"].fill = fill(AZUL); s[f"{col}8"].alignment = Alignment(horizontal="center")
s["A9"] = "Precio sugerido"; s["A10"] = "Lo que te queda neto"; s["A11"] = "Margen real"; s["A12"] = "Comisión + impuestos que pagás"
for col, com in zip("BCD", ["ComClasica+Cuotas", "ComPremium", "ComDirecta"]):
    s[f"{col}9"] = precio("B4+B5", com, "B6", col != "D")
    s[f"{col}10"] = neto(f"{col}9", "B4+B5", com) if col != "D" else f'=IF(ISNUMBER(D9),D9*(1-{com}-IIBB-OtrasRet)-(B4+B5),"")'
    s[f"{col}11"] = f'=IF(ISNUMBER({col}9),{col}10/{col}9,"")'
    s[f"{col}12"] = f'=IF(ISNUMBER({col}9),{col}9-{col}10-(B4+B5),"")'
    for r, fmt in ((9, ARS), (10, ARS), (11, PCT), (12, ARS)):
        s[f"{col}{r}"].number_format = fmt; s[f"{col}{r}"].fill = fill(CLARO); s[f"{col}{r}"].border = box
for r in range(9, 13): s[f"A{r}"].font = B
s["A14"] = "En venta directa no se aplica el costo fijo de Mercado Libre."; s["A14"].font = Font(italic=True, color="595959")

for w in wb.worksheets: w.sheet_view.showGridLines = False
wb.active = 0
wb.save(OUT); print("ok", OUT)
