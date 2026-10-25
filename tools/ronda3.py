"""Ronda 3 (HUECOS.md #7-#9): vuelta al cole, SIRADIG y alquiler para estudiantes.
build_site.main() importa PAGES; `python tools/ronda3.py` arma private/ (necesita openpyxl y python-docx; correr local y commitear; NO corre en Cloudflare):
cole.xlsx, siradig.xlsx, pack-alqestud.zip y pack-siradigcombo.zip (ganancias.xlsx + siradig.xlsx)."""
import io, json, os, zipfile
import legal_pages as LP
import gan_data as G
from ronda2 import head, lec, field, js, quien, AV, AVH, PAGO

# ===================================================================== 9. ALQUILER ESTUDIANTES
def of_est():
    return LP.oferta("alqestud", "Pack en Word (.docx)", "Pack Alquiler para Estudiantes",
        "Los papeles para alquilar una habitación o un departamento compartido para estudiar, con la garantía de un familiar.",
        ["Contrato de locación de habitación en vivienda compartida", "Reglamento de convivencia entre compañeros", "Fianza de un familiar (garante) y recibo de alquiler", "Inventario del cuarto y guía: qué pedir y qué firmar"],
        "2.999", "Comprar el pack", PAGO, '<a class="nota" href="/contratos-alquiler-modelo-word">¿Alquilás una vivienda entera? Pack de contratos de alquiler</a>')

CHK_E = ["Quien alquila es el titular de la vivienda o tiene autorización escrita del dueño para subalquilar: pedí el DNI y un comprobante de titularidad o administración.",
         "Viste la habitación y la casa en persona o en videollamada con alguien de confianza, no solo por fotos.",
         "Quedó escrito el precio, qué servicios incluye (luz, gas, internet) y cómo se reparten las expensas.",
         "Sabés quiénes viven en la casa y las reglas de convivencia (visitas, ruido, limpieza, mascotas) están por escrito.",
         "El plazo del contrato y la forma de salir antes (aviso y penalidad) están claros.",
         "El depósito en garantía figura en un recibo, con la condición para devolverlo.",
         "Firmaste un contrato y no solo una reserva por mensaje, y recibiste copia.",
         "Hay un garante o fianza definida de antemano: quién es y qué cubre.",
         "Hiciste un inventario del cuarto, con fotos fechadas, el día de la entrega."]

BODY_E = (head("Gratis · Alquiler para estudiantes", "Alquilar una habitación para estudiar: qué firmar y cuánto sale", "Antes de señar, tildá lo que pudiste verificar y calculá el costo real por persona, con expensas, servicios y depósito.")
 + '<div class="r2dos"><div><h2 style="margin-top:0">Checklist antes de señar</h2><ul class="r2chk" id="es-chk">' + "".join(f'<li><label><input type="checkbox"> <span>{c}</span></label></li>' for c in CHK_E)
 + '</ul><div class="r2bar"><i id="es-bar"></i></div><p class="nota" id="es-sc"></p></div>'
 + '<div class="comp"><div class="comp-cab"><span>Costo mensual y para ingresar</span><span>Gratis</span></div><div class="comp-cuerpo">'
 + '<div class="fila2">' + field("es-alq", "Alquiler de la habitación", 220000, "$") + field("es-exp", "Expensas (tu parte)", 30000, "$") + "</div>"
 + '<div class="fila2">' + field("es-serv", "Luz, gas e internet (tu parte)", 35000, "$") + field("es-dep", "Meses de depósito", 1, "", "", "1") + "</div>"
 + '<div class="r2res" aria-live="polite"><div class="fil"><span>Costo mensual total</span><span id="es-m"></span></div><div class="fil"><span>Depósito en garantía</span><span id="es-d"></span></div><div class="fil"><span>Para ingresar (1.er mes + depósito)</span><span id="es-i"></span></div><div class="fil"><span>Costo del año de cursada (10 meses)</span><span id="es-a"></span></div></div>'
 + '<p class="nota">El precio, el depósito y los ajustes se acuerdan libremente: lo que vale es lo que quede escrito en el contrato.</p></div></div></div>'
 + AVH
 + lec("<h2>Qué papel hace falta y quién interviene</h2>") + quien([
    ("Contrato por escrito", "Para una habitación en vivienda compartida alcanza con un contrato firmado por las partes, con recibos de cada pago y un inventario. No hace falta escribano.", False),
    ("Garantía", "Lo más común es la fianza de un familiar con inmueble o recibo de sueldo. Acordá con el dueño qué garantía acepta antes de firmar.", False),
    ("Cuándo sí ver a un profesional", "Si el contrato lo firma alguien que no es titular, o si el depósito o la penalidad son altos, pedí a un abogado o escribano que lo revise.", True)])
 + lec("<h2>Qué suma el pack</h2>")
 + LP.modelos_html([dict(titulo="Contrato de habitación", cuando="Para alquilar un cuarto en una casa compartida."), dict(titulo="Reglamento de convivencia", cuando="Visitas, limpieza, ruido y gastos comunes."),
                    dict(titulo="Fianza del garante", cuando="Un familiar se obliga a responder por el alquiler."), dict(titulo="Inventario del cuarto", cuando="Con fotos, al ingresar y al salir."),
                    dict(titulo="Recibo de alquiler", cuando="Cada pago, con fecha y concepto."), dict(titulo="Guía: qué pedir y qué firmar", cuando="Paso a paso, antes de señar.")])
 + of_est())

JS_E = js("""const c=document.querySelectorAll('#es-chk input');
function a(){const m=n('es-alq')+n('es-exp')+n('es-serv'),d=n('es-alq')*n('es-dep');$('es-m').textContent=ars(m);$('es-d').textContent=ars(d);$('es-i').textContent=ars(m+d);$('es-a').textContent=ars(m*10)}
function k(){const x=[...c].filter(e=>e.checked).length;$('es-bar').style.width=(x*100/c.length)+'%';$('es-sc').textContent=x+' de '+c.length+' verificados. '+(x<6?'Todavía no señes: faltan controles.':x<c.length?'Vas bien: completá lo que falta antes de pagar.':'Verificaste todo lo que se puede verificar.')}
document.querySelectorAll('.comp input').forEach(e=>e.addEventListener('input',a));c.forEach(e=>e.addEventListener('change',k));a();k();""")

# ===================================================================== 7. VUELTA AL COLE
NIVELES = {"i": ("Inicial", [35000, 45000, 60000, 0]), "p": ("Primaria", [60000, 70000, 80000, 50000]), "s": ("Secundaria", [85000, 90000, 90000, 120000])}
RUB_C = ["Útiles y mochila", "Uniforme y ropa", "Calzado", "Libros"]

def of_cole():
    return LP.oferta("cole", "Planilla (Excel / Google Sheets)", "Planilla Anual del Gasto Escolar",
        "Presupuestá el año completo por hijo, con la cuota ajustada por IPC, y compará lo presupuestado con lo que gastás de verdad.",
        ["Hasta 3 hijos: útiles, uniforme, cuota, transporte, comedor, libros y actividades", "Cuota y transporte con ajuste mensual por IPC que cargás vos", "Presupuesto vs. gasto real, mes por mes", "Resumen familiar con totales y diferencia"],
        "2.999", "Comprar la planilla", "Pagás con Mercado Pago y la descargás al instante. Funciona en Excel y en Google Sheets.",
        '<a class="nota" href="/cuota-alimentaria-aguinaldo">¿Compartís gastos con el otro progenitor? Documentos de cuota alimentaria</a>')

BODY_C = (head("Gratis · Vuelta al cole", "¿Cuánto me sale la vuelta al cole? Calculá la mochila", "Elegí el nivel, editá los precios con los de tu lista y fijá cuántos hijos tenés. Comparamos pagar de contado y en cuotas.")
 + '<div class="r2dos"><div class="comp"><div class="comp-cab"><span>Tu lista</span><span>Gratis</span></div><div class="comp-cuerpo">'
 + '<div class="campo"><label for="co-n">Nivel</label><select id="co-n">' + "".join(f'<option value="{k}"{" selected" if k=="p" else ""}>{v[0]}</option>' for k, v in NIVELES.items()) + '</select></div>'
 + "".join(f'<div class="fila2">{field(f"co-{i}", r, NIVELES["p"][1][i], "$")}</div>' for i, r in enumerate(RUB_C))
 + '<div class="fila2">' + field("co-h", "Hijos en esa lista", 1, "", "", "1") + field("co-q", "Cuotas sin interés / con recargo", 3, "", "", "1") + "</div>"
 + '<div class="fila2">' + field("co-r", "Recargo por cuotas", 15, "", "%") + field("co-dc", "Descuento por contado", 10, "", "%") + "</div>"
 + '</div></div><div><div class="r2res" id="co-out" aria-live="polite"></div>'
 + '<p class="nota">Los precios iniciales son una referencia: reemplazalos por los de tu lista y tu comercio. No son precios oficiales.</p></div></div>'
 + AVH.replace("Modelos orientativos.", "Cálculo orientativo.").replace(AV, "Valores de referencia que editás vos: no son precios de ningún comercio ni de ningún colegio.")
 + lec("<h2>Cómo bajar el gasto sin quedarte sin lo necesario</h2><ul><li><b>Revisá lo que ya tenés</b> antes de comprar: mochila, calzado y útiles del año anterior suelen servir.</li><li><b>Compará el precio por unidad</b>, no por pack, y fijate si el comercio escolar tiene la lista completa en un solo precio.</li><li><b>Comprá lo no urgente en marzo</b>: libros y materiales que pide el docente tras la primera semana.</li><li><b>Contado o cuotas:</b> si el recargo por cuotas es mayor que la inflación mensual esperada, conviene el contado.</li></ul>")
 + lec("<h2>Qué suma la planilla</h2>")
 + LP.modelos_html([dict(titulo="Una hoja por hijo (hasta 3)", cuando="Útiles, uniforme, calzado, libros, cuota, transporte, comedor y actividades."), dict(titulo="Ajuste por IPC", cuando="Cuota y transporte suben mes a mes con el porcentaje que cargues."),
                    dict(titulo="Presupuesto vs real", cuando="Cargás lo gastado cada mes y ves la diferencia."), dict(titulo="Resumen familiar", cuando="Total anual, por hijo y por mes.")])
 + of_cole())

JS_C = js("const N=" + json.dumps({k: v[1] for k, v in NIVELES.items()}) + """;
function lv(){const v=N[$('co-n').value];v.forEach((x,i)=>$('co-'+i).value=x);a()}
function a(){const t=[0,1,2,3].reduce((s,i)=>s+n('co-'+i),0)*Math.max(1,n('co-h')),q=Math.max(1,n('co-q')),re=n('co-r')/100,dc=n('co-dc')/100;
$('co-out').innerHTML='<div class="fil"><span>Total de la lista</span><span>'+ars(t)+'</span></div><div class="fil"><span>De contado (con descuento)</span><span>'+ars(t*(1-dc))+'</span></div><div class="fil"><span>En '+q+' cuotas (con recargo)</span><span>'+ars(t*(1+re))+'</span></div><div class="fil"><span>Cada cuota</span><span>'+ars(t*(1+re)/q)+'</span></div><p class="nota" style="margin:0">Pagando de contado ahorrás '+ars(t*(1+re)-t*(1-dc))+' frente a las cuotas.</p>'}
$('co-n').addEventListener('change',lv);document.querySelectorAll('.comp input').forEach(e=>e.addEventListener('input',a));a();""")

# ===================================================================== 8. SIRADIG
def of_sir():
    return LP.oferta("siradig", "Planilla (Excel / Google Sheets)", "Organizador SIRADIG: comprobantes y deducciones",
        "Juntá durante el año todos los comprobantes que podés cargar en el SIRADIG y calculá cuánto deducís y cuánto ahorrás en Ganancias.",
        ["Registro de comprobantes con rubro, CUIT, número y monto", "Tope de cada deducción aplicado solo, con las tablas 2026 de ARCA", "Estimación del ahorro de Ganancias según tu sueldo", "Guía de carga paso a paso y marca de «cargado en SIRADIG»"],
        "3.999", "Comprar el organizador", "Pagás con Mercado Pago y lo descargás al instante. Funciona en Excel y en Google Sheets.",
        '<a class="nota" href="/comprar?p=siradigcombo">Combo con la Planilla de Ganancias en el Sueldo: <b>$ 7.999</b> en vez de $ 9.998</a>')

DED_SIR = [("Alquiler de tu vivienda", "40% de lo pagado, hasta la ganancia no imponible."), ("Servicio doméstico", "Lo pagado, hasta la ganancia no imponible."),
           ("Gastos médicos", "40% de lo facturado, hasta el 5% de tu ganancia neta."), ("Prepaga y seguro médico", "Hasta el 5% de tu ganancia neta."),
           ("Donaciones", "Hasta el 5% de tu ganancia neta."), ("Intereses de crédito hipotecario", "Tope anual a verificar en ARCA."), ("Seguro de vida", "Tope anual a verificar en ARCA.")]
BODY_S = (head("Gratis · SIRADIG · Formulario 572", "¿Qué puedo deducir en el SIRADIG? Estimá tu ahorro", "Cargá lo que gastaste en el año y te mostramos cuánto podés deducir según el tope de cada gasto y cuánto menos de Ganancias te retienen.")
 + '<div class="r2dos"><div class="comp"><div class="comp-cab"><span>Tus gastos del año</span><span>Gratis</span></div><div class="comp-cuerpo">'
 + field("sd-neta", "Ganancia neta anual estimada (sueldo del año menos descuentos)", 30000000, "$")
 + '<div class="fila2">' + field("sd-alq", "Alquiler pagado", 6000000, "$") + field("sd-dom", "Servicio doméstico", 0, "$") + "</div>"
 + '<div class="fila2">' + field("sd-med", "Gastos médicos facturados", 800000, "$") + field("sd-pre", "Prepaga y seguro médico", 1500000, "$") + "</div>"
 + '<div class="fila2">' + field("sd-don", "Donaciones", 0, "$") + field("sd-hi", "Hijos a cargo", 1, "", "", "1") + "</div>"
 + '</div></div><div><div class="r2res" id="sd-out" aria-live="polite"></div>'
 + '<p class="nota">Estimación con la escala anual 2026 de ARCA a diciembre. Intereses hipotecarios y seguro de vida tienen topes propios: verificalos en ARCA.</p></div></div>'
 + AVH.replace("Modelos orientativos.", "Cálculo orientativo.").replace(AV, "No reemplaza al contador ni al cálculo de tu empleador. Verificá topes y requisitos vigentes en ARCA.")
 + lec("<h2>Qué se puede cargar y con qué tope</h2><ul>" + "".join(f"<li><b>{t}:</b> {x}</li>" for t, x in DED_SIR) + "</ul>")
 + lec("<h2>Cómo se carga</h2><ol><li>Entrá a arca.gob.ar con Clave Fiscal y abrí SIRADIG - Trabajador.</li><li>Creá o editá el formulario 572 web del período.</li><li>Cargá cada deducción con el CUIT del proveedor y el comprobante.</li><li>Enviá el formulario: tu empleador lo toma para ajustar la retención.</li></ol><p>Tené los comprobantes ordenados antes de que cierre el plazo (la fecha de cierre es en marzo: confirmala en ARCA). Para ver cuánto te cambia cada gasto, probá la <a href=\"/calculadora-ganancias-sueldo\">calculadora de Ganancias</a>.</p>")
 + lec("<h2>Qué suma el organizador</h2>")
 + LP.modelos_html([dict(titulo="Registro de comprobantes", cuando="Fecha, rubro, CUIT, número, monto y marca «cargado»."), dict(titulo="Topes aplicados", cuando="40%, 5% de la neta y ganancia no imponible, solos."),
                    dict(titulo="Ahorro estimado", cuando="Con la escala 2026 de ARCA y tu ganancia neta."), dict(titulo="Guía de carga", cuando="Paso a paso en SIRADIG.")])
 + of_sir())

_esc = G.escala(12); _ded = G.deducciones(12)
JS_S = js("const E=" + json.dumps([[d, f, p] for d, f, p in _esc]) + ",GNI=" + json.dumps(_ded["gni"]) + ",HIJO=" + json.dumps(_ded["hijo"]) + """;
const imp=x=>{if(x<=0)return 0;let t=E[0];for(const r of E)if(r[0]<=x)t=r;return t[1]+(x-t[0])*t[2]/100};
function a(){const neta=n('sd-neta'),al=Math.min(n('sd-alq')*.4,GNI),dm=Math.min(n('sd-dom'),GNI),m5=neta*.05,me=Math.min(n('sd-med')*.4,m5),pr=Math.min(n('sd-pre'),m5),dn=Math.min(n('sd-don'),m5),hi=n('sd-hi')*HIJO,tot=al+dm+me+pr+dn+hi;
const ah=imp(neta-GNI)-imp(neta-GNI-tot);
const f=(l,v)=>'<div class="fil"><span>'+l+'</span><span>'+ars(v)+'</span></div>';
$('sd-out').innerHTML=f('Alquiler deducible',al)+f('Servicio doméstico',dm)+f('Gastos médicos',me)+f('Prepaga y seguro',pr)+f('Donaciones',dn)+f('Hijos',hi)+f('Total deducible',tot)+'<div class="fil"><span><b>Ahorro de Ganancias estimado</b></span><span>'+ars(ah)+'</span></div>'}
document.querySelectorAll('.comp input').forEach(e=>e.addEventListener('input',a));a();""")

PAGES = {
 "contrato-alquiler-estudiante": ("Alquiler para estudiantes: contrato de habitación y checklist", "Checklist gratis y calculadora de costo por persona para alquilar una habitación para estudiar. Pack con contrato de habitación, reglamento y fianza del garante por $2.999.", BODY_E, JS_E),
 "presupuesto-vuelta-al-cole": ("Vuelta al cole: calculá cuánto sale la lista de útiles", "Calculadora gratis del gasto escolar por nivel y por hijo, de contado o en cuotas. Planilla anual de gastos escolares con ajuste por IPC por $2.999.", BODY_C, JS_C),
 "siradig-que-puedo-deducir": ("SIRADIG: qué podés deducir de Ganancias y cuánto ahorrás", "Calculadora gratis de deducciones del SIRADIG con topes 2026: alquiler, médicos, prepaga, donaciones e hijos. Organizador de comprobantes por $3.999.", BODY_S, JS_S),
}

# ===================================================================== archivos
def build_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    priv = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "private")
    AZUL = "1F4E79"; INP = "FFF2CC"; CAL = "DDEBF7"
    H = Font(bold=True, color="FFFFFF"); fill = lambda c: PatternFill("solid", fgColor=c)
    thin = Side(style="thin", color="BFBFBF"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
    ARS = '"$" #,##0'

    def titulo(ws, t, s):
        ws["A1"] = t; ws["A1"].font = Font(bold=True, size=16, color=AZUL); ws["A2"] = s; ws["A2"].font = Font(italic=True, color="595959")
    def cab(ws, row, heads, widths=None):
        for i, t in enumerate(heads, 1):
            c = ws.cell(row=row, column=i, value=t); c.font = H; c.fill = fill(AZUL); c.border = box; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if widths: ws.column_dimensions[c.column_letter].width = widths[i - 1]
    def inicio(ws, lines):
        ws.column_dimensions["A"].width = 120
        for i, l in enumerate(lines, 4):
            ws.cell(row=i, column=1, value=l)
            if l.isupper() and l: ws.cell(row=i, column=1).font = Font(bold=True, color=AZUL, size=12)
    def st(c, kind, fmt=None):
        c.fill = fill(INP if kind == "i" else CAL); c.border = box
        if fmt: c.number_format = fmt

    # ---------------- cole.xlsx
    wb = Workbook(); ws = wb.active; ws.title = "Inicio"
    titulo(ws, "Planilla Anual del Gasto Escolar", "Presupuesto por hijo, ajuste por IPC y comparación con el gasto real. Excel y Google Sheets.")
    inicio(ws, ["CÓMO USARLO", "1. Pestaña 'Parámetros': poné el IPC mensual que esperás (por ejemplo 3%) y los nombres de tus hijos.",
        "2. En cada pestaña 'Hijo': en 'Gastos mensuales' cargá el monto base, si ajusta por IPC (SI/NO) y de qué mes a qué mes se paga. El resto se calcula.",
        "3. En 'Gastos únicos' cargá útiles, uniforme, calzado, libros y otros en el mes en que los vas a comprar.",
        "4. En 'Gastado real' cargá el total que gastaste cada mes. La diferencia contra lo presupuestado se calcula sola.",
        "5. 'Resumen' suma a toda la familia, por mes y por hijo.", "", "Celdas amarillas = las cargás vos. Celdas celestes = se calculan solas. Los números de ejemplo son referencia: pisalos con los tuyos.",
        "", "AVISO", "Herramienta de cálculo orientativa. Los precios de ejemplo no corresponden a ningún colegio ni comercio. Para uso propio, prohibida su reventa."])
    MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
    wp = wb.create_sheet("Parámetros"); titulo(wp, "Parámetros", "Valores que usan todas las pestañas")
    wp["A4"] = "IPC mensual esperado"; wp["B4"] = 0.03; st(wp["B4"], "i", "0.0%")
    wp.column_dimensions["A"].width = 28; wp.column_dimensions["B"].width = 18
    for i in range(3):
        wp.cell(row=6 + i, column=1, value=f"Nombre hijo {i+1}"); wp.cell(row=6 + i, column=2, value=f"Hijo {i+1}"); st(wp.cell(row=6 + i, column=2), "i")
    for h in range(1, 4):
        w = wb.create_sheet(f"Hijo {h}"); titulo(w, f"Hijo {h}", "Presupuesto y gasto real del año")
        w["A3"] = "Nombre"; w["B3"] = f"=Parámetros!B{5+h}"; w["B3"].font = Font(bold=True)
        cab(w, 5, ["Gastos mensuales", "Monto base", "Ajusta IPC", "Desde (mes 1-12)", "Hasta (mes 1-12)"] + MESES + ["Total año"], [34, 14, 12, 14, 14] + [12] * 12 + [16])
        mens = [("Cuota del colegio", 120000 if h == 1 else 0, "SI", 3, 12), ("Transporte", 45000 if h == 1 else 0, "SI", 3, 12), ("Comedor", 0, "NO", 3, 12), ("Actividades extra", 0, "NO", 3, 11)]
        dv = DataValidation(type="list", formula1='"SI,NO"', allow_blank=False); w.add_data_validation(dv)
        for r, (nom, base, aj, d, hh) in enumerate(mens, 6):
            w.cell(row=r, column=1, value=nom);
            for col, v, f in ((2, base, ARS), (3, aj, None), (4, d, "0"), (5, hh, "0")):
                c = w.cell(row=r, column=col, value=v); st(c, "i", f)
            dv.add(w.cell(row=r, column=3))
            for m in range(1, 13):
                col = 5 + m; cl = w.cell(row=r, column=col)
                cl.value = f'=IF(AND({m}>=$D{r},{m}<=$E{r}),$B{r}*IF($C{r}="SI",(1+Parámetros!$B$4)^({m}-$D{r}),1),0)'; st(cl, "c", ARS)
            t = w.cell(row=r, column=18, value=f"=SUM(F{r}:Q{r})"); st(t, "c", ARS)
        cab(w, 11, ["Gastos únicos (cargalos en su mes)", "", "", "", ""] + MESES + ["Total año"])
        unic = [("Útiles y mochila", {2: 60000}), ("Uniforme y ropa", {2: 70000}), ("Calzado", {2: 80000}), ("Libros", {3: 50000}), ("Otros", {})]
        for r, (nom, ej) in enumerate(unic, 12):
            w.cell(row=r, column=1, value=nom)
            for m in range(1, 13):
                c = w.cell(row=r, column=5 + m, value=ej.get(m, 0) if h == 1 else 0); st(c, "i", ARS)
            t = w.cell(row=r, column=18, value=f"=SUM(F{r}:Q{r})"); st(t, "c", ARS)
        w.cell(row=18, column=1, value="TOTAL PRESUPUESTADO").font = Font(bold=True)
        w.cell(row=19, column=1, value="Gastado real").font = Font(bold=True)
        w.cell(row=20, column=1, value="Diferencia (presupuesto - real)").font = Font(bold=True)
        for col in range(6, 19):
            L = w.cell(row=5, column=col).column_letter
            c = w.cell(row=18, column=col, value=f"=SUM({L}6:{L}9)+SUM({L}12:{L}16)"); st(c, "c", ARS)
            if col < 18: st(w.cell(row=19, column=col), "i", ARS)
            else: w.cell(row=19, column=col, value="=SUM(F19:Q19)"); st(w.cell(row=19, column=col), "c", ARS)
            c = w.cell(row=20, column=col, value=f"={L}18-{L}19"); st(c, "c", ARS)
        w.freeze_panes = "B6"
    wr = wb.create_sheet("Resumen"); titulo(wr, "Resumen familiar", "Suma de los tres hijos")
    cab(wr, 4, ["Concepto"] + MESES + ["Total año"], [32] + [13] * 12 + [16])
    for r, (nom, row) in enumerate([("Presupuestado", 18), ("Gastado real", 19), ("Diferencia", 20)], 5):
        wr.cell(row=r, column=1, value=nom).font = Font(bold=True)
        for col in range(2, 15):
            L = wr.cell(row=4, column=col).column_letter; src = wr.cell(row=4, column=col).column_letter
            sc = chr(ord("F") + col - 2) if col < 14 else "R"
            f = "=" + "+".join(f"'Hijo {h}'!{sc}{row}" for h in (1, 2, 3))
            st(wr.cell(row=r, column=col, value=f), "c", ARS)
    wb.save(os.path.join(priv, "cole.xlsx")); print("ok cole.xlsx")

    # ---------------- siradig.xlsx
    wb = Workbook(); ws = wb.active; ws.title = "Inicio"
    titulo(ws, "Organizador SIRADIG 2026", "Registro de comprobantes deducibles, topes y ahorro estimado de Ganancias. Excel y Google Sheets.")
    inicio(ws, ["CÓMO USARLO", "1. Pestaña 'Comprobantes': cargá cada factura o recibo (fecha, rubro, CUIT del proveedor, número, monto). Marcá SI cuando lo cargues en el SIRADIG.",
        "2. Pestaña 'Resumen': poné tu ganancia neta anual estimada y la cantidad de hijos. Te muestra cuánto podés deducir por rubro con su tope y el ahorro estimado.",
        "3. Pestaña 'Parámetros': escala y deducciones personales de ARCA 2026 (acumulado a diciembre) y los topes que debés completar (hipotecario, seguro de vida).",
        "4. Pestaña 'Guía de carga': los pasos para cargar el formulario 572 web.", "",
        "Celdas amarillas = las cargás vos. Celdas celestes = se calculan solas. Las filas de ejemplo se pisan o se borran.", "",
        "TOPES QUE APLICA", "Alquiler: 40% de lo pagado, hasta la ganancia no imponible. Servicio doméstico: lo pagado, hasta la ganancia no imponible.",
        "Gastos médicos: 40% de lo facturado, hasta el 5% de la ganancia neta. Prepaga y seguro médico: hasta el 5% de la ganancia neta. Donaciones: hasta el 5% de la ganancia neta.",
        "Intereses hipotecarios y seguro de vida: completá el tope anual en 'Parámetros' (si queda vacío, no se aplica tope y la hoja avisa).", "",
        "AVISO", "Estimación orientativa. No reemplaza el cálculo de tu empleador ni el de un contador. Verificá topes, requisitos y la fecha de cierre del SIRADIG en ARCA.", "Para uso propio, prohibida su reventa."])
    wp = wb.create_sheet("Parámetros"); titulo(wp, "Parámetros 2026", "Escala acumulada a diciembre y deducciones personales (fuente: ARCA, art. 30 y 94)")
    wp.column_dimensions["A"].width = 38
    for c in "BCD": wp.column_dimensions[c].width = 18
    ded = [("Ganancia no imponible", _ded["gni"]), ("Cónyuge", _ded["conyuge"]), ("Hijo", _ded["hijo"])]
    for r, (nom, v) in enumerate(ded, 4):
        wp.cell(row=r, column=1, value=nom); st(wp.cell(row=r, column=2, value=round(v, 2)), "i", ARS)
    wp["A8"] = "Tope anual intereses hipotecarios"; wp["A9"] = "Tope anual seguro de vida"
    for r in (8, 9): st(wp.cell(row=r, column=2), "i", ARS)
    cab(wp, 11, ["Desde (neta anual)", "Monto fijo", "% sobre excedente"])
    for r, (d, f, p) in enumerate(_esc, 12):
        st(wp.cell(row=r, column=1, value=d), "i", ARS); st(wp.cell(row=r, column=2, value=f), "i", ARS); st(wp.cell(row=r, column=3, value=p), "i", "0")
    wc = wb.create_sheet("Comprobantes"); titulo(wc, "Comprobantes", "Un renglón por factura o recibo")
    RUB = ["Alquiler", "Servicio doméstico", "Gastos médicos", "Prepaga y seguro médico", "Donaciones", "Intereses hipotecarios", "Seguro de vida"]
    cab(wc, 4, ["Fecha", "Rubro", "CUIT del proveedor", "N.º de comprobante", "Descripción", "Monto", "Cargado en SIRADIG"], [13, 28, 20, 20, 38, 16, 18])
    dv = DataValidation(type="list", formula1='"' + ",".join(RUB) + '"'); dv2 = DataValidation(type="list", formula1='"SI,NO"')
    wc.add_data_validation(dv); wc.add_data_validation(dv2)
    ej = [("2026-03-05", "Alquiler", "Alquiler marzo", 500000), ("2026-04-05", "Alquiler", "Alquiler abril", 500000), ("2026-03-12", "Gastos médicos", "Consulta y estudios", 85000)]
    import datetime
    for r in range(5, 205):
        for col in range(1, 8): wc.cell(row=r, column=col).border = box
        wc.cell(row=r, column=1).number_format = "dd/mm/yyyy"; wc.cell(row=r, column=6).number_format = ARS
        for col in (1, 2, 3, 4, 5, 6, 7): wc.cell(row=r, column=col).fill = fill(INP)
        dv.add(wc.cell(row=r, column=2)); dv2.add(wc.cell(row=r, column=7))
    for r, (f, ru, de, mo) in enumerate(ej, 5):
        wc.cell(row=r, column=1, value=datetime.date.fromisoformat(f)); wc.cell(row=r, column=2, value=ru); wc.cell(row=r, column=5, value=de); wc.cell(row=r, column=6, value=mo); wc.cell(row=r, column=7, value="NO")
    wc.freeze_panes = "A5"
    wr = wb.create_sheet("Resumen"); titulo(wr, "Resumen de deducciones", "Cuánto podés deducir por rubro y cuánto ahorrás")
    wr.column_dimensions["A"].width = 38
    for c in "BCDE": wr.column_dimensions[c].width = 22
    wr["A4"] = "Ganancia neta anual estimada"; st(wr["B4"], "i", ARS); wr["B4"] = 30000000
    wr["A5"] = "Hijos a cargo"; st(wr["B5"], "i", "0"); wr["B5"] = 1
    wr["A6"] = "Cónyuge a cargo (1 = sí, 0 = no)"; st(wr["B6"], "i", "0"); wr["B6"] = 0
    cab(wr, 8, ["Rubro", "Total cargado", "Tope aplicable", "Deducible", "Aviso"])
    topes = {"Alquiler": "=Parámetros!$B$4", "Servicio doméstico": "=Parámetros!$B$4", "Gastos médicos": "=0.05*$B$4", "Prepaga y seguro médico": "=0.05*$B$4", "Donaciones": "=0.05*$B$4",
             "Intereses hipotecarios": '=IF(Parámetros!$B$8="","sin tope cargado",Parámetros!$B$8)', "Seguro de vida": '=IF(Parámetros!$B$9="","sin tope cargado",Parámetros!$B$9)'}
    pct = {"Alquiler": 0.4, "Gastos médicos": 0.4}
    for r, ru in enumerate(RUB, 9):
        wr.cell(row=r, column=1, value=ru); st(wr.cell(row=r, column=2, value=f'=SUMIFS(Comprobantes!$F$5:$F$204,Comprobantes!$B$5:$B$204,A{r})'), "c", ARS)
        st(wr.cell(row=r, column=3, value=topes[ru]), "c", ARS)
        base = f"B{r}*{pct[ru]}" if ru in pct else f"B{r}"
        st(wr.cell(row=r, column=4, value=f"=IF(ISNUMBER(C{r}),MIN({base},C{r}),{base})"), "c", ARS)
        if ru in ("Intereses hipotecarios", "Seguro de vida"):
            wr.cell(row=r, column=5, value=f'=IF(ISNUMBER(C{r}),"","Completá el tope en Parámetros")')
    wr["A16"] = "Hijos"; st(wr["D16"], "c", ARS); wr["D16"] = "=B5*Parámetros!$B$6"
    wr["A17"] = "Cónyuge"; st(wr["D17"], "c", ARS); wr["D17"] = "=B6*Parámetros!$B$5"
    wr["A18"] = "TOTAL DEDUCIBLE"; wr["A18"].font = Font(bold=True); st(wr["D18"], "c", ARS); wr["D18"] = "=SUM(D9:D17)"
    imp = lambda x: f"(LOOKUP({x},Parámetros!$A$12:$A$20,Parámetros!$B$12:$B$20)+({x}-LOOKUP({x},Parámetros!$A$12:$A$20,Parámetros!$A$12:$A$20))*LOOKUP({x},Parámetros!$A$12:$A$20,Parámetros!$C$12:$C$20)/100)"
    wr["A20"] = "Neta tras ganancia no imponible"; st(wr["D20"], "c", ARS); wr["D20"] = "=MAX(0,B4-Parámetros!B4)"
    wr["A21"] = "Impuesto sin deducciones"; st(wr["D21"], "c", ARS); wr["D21"] = "=" + imp("D20")
    wr["A22"] = "Impuesto con deducciones"; st(wr["D22"], "c", ARS); wr["D22"] = "=" + imp("MAX(0,D20-D18)")
    wr["A23"] = "AHORRO ESTIMADO DE GANANCIAS"; wr["A23"].font = Font(bold=True); st(wr["D23"], "c", ARS); wr["D23"] = "=D21-D22"
    wr["A25"] = "Comprobantes cargados / pendientes de cargar en SIRADIG"; wr["D25"] = '=COUNTIF(Comprobantes!G5:G204,"SI")&" / "&COUNTIF(Comprobantes!G5:G204,"NO")'
    wg = wb.create_sheet("Guía de carga"); titulo(wg, "Guía de carga en SIRADIG", "Formulario 572 web de ARCA")
    wg.column_dimensions["A"].width = 120
    for i, l in enumerate(["1. Entrá a arca.gob.ar con tu CUIT y Clave Fiscal.", "2. Abrí el servicio SIRADIG - Trabajador.", "3. Creá el formulario 572 web del período (o editá el que ya tenés).",
            "4. Cargá cada deducción: elegí el tipo (alquiler, gastos médicos, servicio doméstico, etc.), el CUIT del proveedor, el mes, el número de comprobante y el monto.",
            "5. Controlá contra la pestaña 'Comprobantes' que no falte ninguno y marcá SI en 'Cargado en SIRADIG'.", "6. Enviá el formulario: tu empleador lo toma para ajustar la retención.",
            "7. Guardá los comprobantes: ARCA puede pedirlos si se controla la declaración.", "", "El cierre del plazo de carga es en marzo del año siguiente: confirmá la fecha exacta en ARCA.", "Los requisitos de cada deducción (por ejemplo, alquiler declarado, médicos con factura a tu nombre) pueden cambiar: verificalos en ARCA."], 4):
        wg.cell(row=i, column=1, value=l)
    wb.save(os.path.join(priv, "siradig.xlsx")); print("ok siradig.xlsx")

    # ---------------- combo siradig + ganancias
    with zipfile.ZipFile(os.path.join(priv, "pack-siradigcombo.zip"), "w", zipfile.ZIP_DEFLATED) as z:
        z.write(os.path.join(priv, "ganancias.xlsx"), "Planilla-Ganancias-Sueldo-2026.xlsx"); z.write(os.path.join(priv, "siradig.xlsx"), "Organizador-SIRADIG-2026.xlsx")
    print("ok pack-siradigcombo.zip")

def build_docx():
    from packs_lib import Doc
    priv = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "private")
    P = "Pack Alquiler para Estudiantes"
    def mk(t, s, box=None):
        d = Doc(P, t, s); d.box("Antes de usarlo", box or [AV, "Completá los campos resaltados en amarillo y borrá lo que no aplique."]); return d
    d = Doc(P, "Guía y contenido del pack", "Qué pedir y qué firmar")
    d.box("Importante", [AV, "Sirve para el propietario o administrador y para el estudiante: leé el modelo completo antes de firmar."])
    d.h("Qué incluye"); d.bullets(["01. Contrato de locación de habitación", "02. Reglamento de convivencia", "03. Fianza de un familiar (garante)", "04. Inventario del cuarto y recibo (del pack de alquileres)", "05. Checklist antes de señar"])
    d.h("Qué pedir antes de firmar"); d.bullets(["Que quien alquila sea titular o tenga autorización escrita del dueño para subalquilar.", "Contrato por escrito, con recibo de cada pago y copia para vos.", "Que figuren el precio, qué incluye (luz, gas, internet), el plazo y la forma de salir antes.", "Que el depósito conste en un recibo, con la condición para devolverlo.", "Que la garantía (por ejemplo, la fianza de un familiar) esté definida de antemano."])
    d.h("Qué exige escribano o juzgado"); d.bullets(["Un contrato de habitación en vivienda compartida no exige escribano: alcanza con la firma de las partes.", "Algunas provincias cobran impuesto de sellos a los contratos de locación: consultá al organismo recaudador de tu provincia.", "Si el contrato lo firma alguien que no es titular, o el monto o la penalidad son altos, pedí a un abogado o escribano que lo revise."])
    docs = [("00-LEEME-Guia-del-pack.docx", d)]
    d = mk("Contrato de locación de habitación", "Vivienda compartida con fines de estudio")
    d.p("En [[ciudad]], a los [[fecha]], entre [[nombre del locador]], DNI [[dni]], domicilio en [[domicilio]], titular o autorizado para subalquilar el inmueble de [[dirección]] (en adelante, EL LOCADOR), y [[nombre del estudiante]], DNI [[dni]], domicilio [[domicilio]] (en adelante, EL LOCATARIO), convienen:")
    d.clauses([("Objeto", "EL LOCADOR da en locación al LOCATARIO la habitación [[número o descripción]] ubicada en el inmueble indicado, con uso compartido de [[cocina, baño, living, lavadero]]."),
               ("Destino", "La habitación se destina a vivienda del LOCATARIO durante el período de estudios. No podrá subalquilarse ni alojar a terceros de forma permanente."),
               ("Plazo", "El plazo es de [[meses]] meses, desde el [[fecha de inicio]] hasta el [[fecha de fin]]."),
               ("Precio", "El precio mensual es de $ [[monto]], pagadero del [[1]] al [[10]] de cada mes mediante [[transferencia a la cuenta / efectivo con recibo]]. [[Se actualizará cada tantos meses según el índice acordado / no se actualiza durante el plazo]]."),
               ("Servicios y expensas", "El precio [[incluye / no incluye]] luz, gas, agua e internet. Las expensas se pagan [[por el locador / por el locatario / a prorrata entre los ocupantes]]."),
               ("Depósito en garantía", "El LOCATARIO entrega $ [[monto]] en concepto de depósito, que se devuelve dentro de [[días]] días del egreso, una vez controlado el estado de la habitación y descontados daños comprobados."),
               ("Convivencia", "El LOCATARIO acepta el reglamento de convivencia, que forma parte de este contrato."),
               ("Rescisión anticipada", "El LOCATARIO puede dejar la habitación antes del plazo con un aviso de [[días]] días. Si no avisa, [[pierde el depósito / abona un mes de alquiler]]."),
               ("Inventario", "Las partes firman un inventario del estado de la habitación y sus muebles al ingreso y al egreso."),
               ("Garantía", "EL LOCATARIO presenta la garantía de [[nombre del garante]], conforme al documento de fianza adjunto.")])
    d.firmas(["Locador/a", "Locatario/a"]); docs.append(("01-Contrato-habitacion.docx", d))
    d = mk("Reglamento de convivencia", "Normas de la casa compartida")
    d.h("Normas"); d.bullets(["Visitas: hasta las [[hora]]; las visitas que duermen se avisan a los demás con [[anticipación]].", "Horario de silencio: de [[hora]] a [[hora]].", "Limpieza: turnos semanales de [[espacios comunes]] según el cronograma de [[lugar]].", "Gastos comunes (luz, gas, internet, limpieza): se dividen [[por partes iguales / según consumo]] y se pagan antes del [[día]].", "Cocina y heladera: cada uno rotula lo suyo; se lava lo que se usa.", "Mascotas: [[permitidas / no permitidas]]. Fumar: [[permitido en exteriores / prohibido]].", "Llaves: cada ocupante tiene una copia y no la presta; si se pierde, se repone a su cargo.", "Los daños en espacios comunes se reparten entre los ocupantes, salvo que se conozca al responsable."])
    d.p("Contacto del locador: [[nombre y teléfono]].", False); docs.append(("02-Reglamento-de-convivencia.docx", d))
    d = mk("Fianza solidaria", "Garantía de un familiar")
    d.p("En [[ciudad]], a los [[fecha]], [[nombre del fiador]], DNI [[dni]], domicilio en [[domicilio]] y propietario del inmueble de [[dirección y matrícula]] (en adelante, EL FIADOR), declara que se constituye fiador solidario, liso, llano y principal pagador de las obligaciones asumidas por [[nombre del estudiante]], DNI [[dni]], en el contrato de locación de habitación con [[nombre del locador]] de fecha [[fecha]], por el precio de $ [[monto]] mensual y por todo el plazo del contrato.")
    d.clauses([("Alcance", "La fianza cubre el precio, las expensas y servicios a cargo del locatario, los daños comprobados y las costas que se originen por el incumplimiento."),
               ("Renuncias", "EL FIADOR renuncia a los beneficios de excusión y de división."),
               ("Plazo", "La fianza subsiste hasta la devolución de la habitación y el pago de todo lo adeudado."),
               ("Domicilio", "EL FIADOR constituye domicilio en el indicado y acepta las notificaciones que allí se realicen.")])
    d.firmas(["Fiador/a"]); docs.append(("03-Fianza-garante.docx", d))
    d = mk("Checklist antes de señar", "Para estudiantes")
    d.table(["Verificación", "Resultado"], [[c, "[[ ]]"] for c in CHK_E], [11, 4]); docs.append(("05-Checklist-antes-de-senar.docx", d))
    z = zipfile.ZipFile(os.path.join(priv, "pack-alquiler.zip"))
    reuse = [("04-Anexo-Inventario-Estado-Inmueble.docx", z.read("04-Anexo-Inventario-Estado-Inmueble.docx")), ("04b-Recibo-de-Alquiler.docx", z.read("07-Recibo-de-Alquiler.docx"))]
    with zipfile.ZipFile(os.path.join(priv, "pack-alqestud.zip"), "w", zipfile.ZIP_DEFLATED) as zz:
        for fn, dd in docs:
            buf = io.BytesIO(); dd.d.save(buf); zz.writestr(fn, buf.getvalue())
        for fn, data in reuse: zz.writestr(fn, data)
    print("ok pack-alqestud.zip")

if __name__ == "__main__":
    build_xlsx(); build_docx()
