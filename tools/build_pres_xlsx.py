"""Planilla 'Presupuesto familiar 2027 con IPC' (HUECOS.md #20). `python tools/build_pres_xlsx.py` arma private/presupuesto.xlsx (solo local, necesita openpyxl; commitear el .xlsx)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

priv = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "private")
AZUL = "1F4E79"; INP = "FFF2CC"; CAL = "DDEBF7"
H = Font(bold=True, color="FFFFFF"); fill = lambda c: PatternFill("solid", fgColor=c)
thin = Side(style="thin", color="BFBFBF"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
ARS = '"$" #,##0'; MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]

def titulo(ws, t, s):
    ws["A1"] = t; ws["A1"].font = Font(bold=True, size=16, color=AZUL); ws["A2"] = s; ws["A2"].font = Font(italic=True, color="595959")
def cab(ws, row, heads, widths):
    for i, (t, w) in enumerate(zip(heads, widths), 1):
        c = ws.cell(row=row, column=i, value=t); c.font = H; c.fill = fill(AZUL); c.border = box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); ws.column_dimensions[c.column_letter].width = w
def st(c, kind, fmt=None):
    c.fill = fill(INP if kind == "i" else CAL); c.border = box
    if fmt: c.number_format = fmt

wb = Workbook(); ws = wb.active; ws.title = "Inicio"
titulo(ws, "Presupuesto Familiar 2027 con IPC", "Un presupuesto anual en pesos que se ajusta solo por la inflación que cargues. Excel y Google Sheets.")
ws.column_dimensions["A"].width = 120
lines = ["CÓMO USARLO",
 "1. Pestaña 'Parámetros': cargá el IPC esperado para cada mes de 2027 (por ejemplo 2,5%) y tus ingresos de enero con el aumento mensual que esperás.",
 "2. Pestaña 'Presupuesto': en cada rubro poné el gasto base de enero y si ajusta por IPC (SI/NO). Los meses siguientes se calculan solos.",
 "3. Pestaña 'Real': cargá lo que gastaste de verdad cada mes. Compará contra el presupuesto.",
 "4. Pestaña 'Resumen': ingresos, gastos, ahorro, ahorro acumulado y si tus ingresos le ganan a la inflación (poder de compra).",
 "5. Cuando el INDEC publique el IPC real de cada mes, reemplazá el esperado por el real y todo se recalcula (fuente: indec.gob.ar).",
 "", "Celdas amarillas = las cargás vos. Celdas celestes = se calculan solas. Los números de ejemplo son referencia: pisalos con los tuyos.",
 "", "AVISO", "Herramienta de cálculo orientativa. No es asesoramiento financiero. Los valores de ejemplo no son predicciones. Para uso propio, prohibida su reventa."]
for i, l in enumerate(lines, 4):
    ws.cell(row=i, column=1, value=l)
    if l in ("CÓMO USARLO", "AVISO"): ws.cell(row=i, column=1).font = Font(bold=True, color=AZUL, size=12)

wp = wb.create_sheet("Parámetros"); titulo(wp, "Parámetros", "IPC mensual esperado e ingresos")
cab(wp, 4, ["Concepto"] + MESES, [34] + [10] * 12)
wp.cell(row=5, column=1, value="IPC mensual esperado")
for m in range(12): st(wp.cell(row=5, column=2 + m, value=0.025), "i", "0.0%")
wp.cell(row=6, column=1, value="Índice de precios (enero = 1)")
for m in range(12):
    c = wp.cell(row=6, column=2 + m, value=1 if m == 0 else f"={wp.cell(row=6, column=1 + m).column_letter}6*(1+{wp.cell(row=5, column=2 + m).column_letter}5)"); st(c, "c", "0.000")
wp["A8"] = "Ingresos de enero"; wp["B8"] = 1500000; st(wp["B8"], "i", ARS)
wp["A9"] = "Aumento mensual esperado de ingresos"; wp["B9"] = 0.02; st(wp["B9"], "i", "0.0%")
wp["A10"] = "Meta de ahorro mensual (% del ingreso)"; wp["B10"] = 0.1; st(wp["B10"], "i", "0%")
for c in ("A8", "A9", "A10"): wp[c].font = Font(bold=True)

wb2 = wb.create_sheet("Presupuesto"); titulo(wb2, "Presupuesto 2027", "Gasto base de enero y ajuste mensual por IPC")
cab(wb2, 4, ["Rubro", "Base enero", "Ajusta IPC"] + MESES + ["Total año"], [30, 14, 12] + [12] * 12 + [16])
rub = [("Alquiler / cuota de vivienda", 450000, "SI"), ("Expensas y servicios", 120000, "SI"), ("Supermercado", 350000, "SI"), ("Transporte", 80000, "SI"), ("Salud y prepaga", 150000, "SI"),
       ("Educación", 100000, "SI"), ("Ocio y salidas", 80000, "NO"), ("Seguros e impuestos", 40000, "SI"), ("Deudas y cuotas", 60000, "NO"), ("Otros", 30000, "NO")]
dv = DataValidation(type="list", formula1='"SI,NO"', allow_blank=False); wb2.add_data_validation(dv)
for r, (n, b, a) in enumerate(rub, 5):
    wb2.cell(row=r, column=1, value=n); st(wb2.cell(row=r, column=2, value=b), "i", ARS); st(wb2.cell(row=r, column=3, value=a), "i"); dv.add(wb2.cell(row=r, column=3))
    for m in range(12):
        L = wp.cell(row=6, column=2 + m).column_letter
        st(wb2.cell(row=r, column=4 + m, value=f'=$B{r}*IF($C{r}="SI",Parámetros!{L}$6,1)'), "c", ARS)
    st(wb2.cell(row=r, column=16, value=f"=SUM(D{r}:O{r})"), "c", ARS)
tr = 5 + len(rub); wb2.cell(row=tr, column=1, value="TOTAL").font = Font(bold=True)
for col in range(4, 17):
    L = wb2.cell(row=4, column=col).column_letter; st(wb2.cell(row=tr, column=col, value=f"=SUM({L}5:{L}{tr-1})"), "c", ARS)
wb2.freeze_panes = "D5"

wr = wb.create_sheet("Real"); titulo(wr, "Gasto real 2027", "Cargá lo que gastaste cada mes")
cab(wr, 4, ["Rubro", "", ""] + MESES + ["Total año"], [30, 4, 4] + [12] * 12 + [16])
for r, (n, _, _) in enumerate(rub, 5):
    wr.cell(row=r, column=1, value=n)
    for m in range(12): st(wr.cell(row=r, column=4 + m, value=0), "i", ARS)
    st(wr.cell(row=r, column=16, value=f"=SUM(D{r}:O{r})"), "c", ARS)
wr.cell(row=tr, column=1, value="TOTAL").font = Font(bold=True)
for col in range(4, 17):
    L = wr.cell(row=4, column=col).column_letter; st(wr.cell(row=tr, column=col, value=f"=SUM({L}5:{L}{tr-1})"), "c", ARS)

wz = wb.create_sheet("Resumen"); titulo(wz, "Resumen del año", "Ingresos, gastos, ahorro y poder de compra")
cab(wz, 4, ["Concepto", "", ""] + MESES + ["Total año"], [38, 4, 4] + [13] * 12 + [16])
filas = ["Ingresos", "Gasto presupuestado", "Ahorro presupuestado", "Meta de ahorro", "Ahorro acumulado", "Gasto real", "Diferencia (presupuesto - real)", "Ingreso real (a precios de enero)"]
for i, n in enumerate(filas): wz.cell(row=5 + i, column=1, value=n).font = Font(bold=True)
for m in range(12):
    col = 4 + m; L = wz.cell(row=4, column=col).column_letter; P = wp.cell(row=6, column=2 + m).column_letter
    st(wz.cell(row=5, column=col, value=f"=Parámetros!$B$8*(1+Parámetros!$B$9)^{m}"), "c", ARS)
    st(wz.cell(row=6, column=col, value=f"=Presupuesto!{L}{tr}"), "c", ARS)
    st(wz.cell(row=7, column=col, value=f"={L}5-{L}6"), "c", ARS)
    st(wz.cell(row=8, column=col, value=f"={L}5*Parámetros!$B$10"), "c", ARS)
    st(wz.cell(row=9, column=col, value=f"=SUM($D7:{L}7)"), "c", ARS)
    st(wz.cell(row=10, column=col, value=f"=Real!{L}{tr}"), "c", ARS)
    st(wz.cell(row=11, column=col, value=f"={L}6-{L}10"), "c", ARS)
    st(wz.cell(row=12, column=col, value=f"={L}5/Parámetros!{P}6"), "c", ARS)
for r in (5, 6, 7, 8, 10, 11):
    st(wz.cell(row=r, column=16, value=f"=SUM(D{r}:O{r})"), "c", ARS)
wz["A14"] = "Si 'Ingreso real' baja mes a mes, tus ingresos no le ganan a la inflación esperada."; wz["A14"].font = Font(italic=True, color="595959")
wz.freeze_panes = "D5"
wb.save(os.path.join(priv, "presupuesto.xlsx")); print("ok presupuesto.xlsx")
