import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x2A, 0x44); GRAY = RGBColor(0x5B, 0x63, 0x70); ACC = RGBColor(0xB4, 0x4A, 0x1E)
BODY = "Cambria"; HEAD = "Calibri"
ORD = ["PRIMERA","SEGUNDA","TERCERA","CUARTA","QUINTA","SEXTA","SÉPTIMA","OCTAVA","NOVENA","DÉCIMA","DÉCIMA PRIMERA","DÉCIMA SEGUNDA","DÉCIMA TERCERA","DÉCIMA CUARTA","DÉCIMA QUINTA","DÉCIMA SEXTA","DÉCIMA SÉPTIMA","DÉCIMA OCTAVA","DÉCIMA NOVENA","VIGÉSIMA"]
AVISO = "Modelo orientativo. No reemplaza el asesoramiento de un abogado/a. Revisalo antes de firmar."

def _font(run, name=BODY, size=10.5, bold=None, italic=None, color=None):
    run.font.name = name; run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color is not None: run.font.color.rgb = color

def _shade(el, fill):
    pr = el.get_or_add_tcPr() if hasattr(el, "get_or_add_tcPr") else el
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill); pr.append(sh)

def _border(par, side="bottom", color="1F2A44", sz=8):
    pPr = par._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); e = OxmlElement(f"w:{side}")
    for k, v in (("val","single"),("sz",str(sz)),("space","4"),("color",color)): e.set(qn(f"w:{k}"), v)
    b.append(e); pPr.append(b)

def _field(par, instr):
    r = par.add_run(); _font(r, HEAD, 8, color=GRAY)
    for t, txt in (("begin", None), (None, instr), ("end", None)):
        if t:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t); r._r.append(f)
        else:
            i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = txt; r._r.append(i)

def rich(par, text, size=10.5, color=None, bold=None, italic=None, name=BODY):
    """**negrita**, [[campo a completar]] resaltado."""
    for part in re.split(r"(\*\*.+?\*\*|\[\[.+?\]\])", text):
        if not part: continue
        if part.startswith("**"):
            r = par.add_run(part[2:-2]); _font(r, name, size, True, italic, color)
        elif part.startswith("[["):
            r = par.add_run("[" + part[2:-2] + "]"); _font(r, name, size, True, italic, NAVY)
            rpr = r._r.get_or_add_rPr(); h = OxmlElement("w:shd"); h.set(qn("w:val"),"clear"); h.set(qn("w:color"),"auto"); h.set(qn("w:fill"),"FFF1C2"); rpr.append(h)
        else:
            r = par.add_run(part); _font(r, name, size, bold, italic, color)

class Doc:
    def __init__(self, pack, titulo, sub=None):
        self.d = Document(); s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.5); s.top_margin = Cm(2.3); s.bottom_margin = Cm(2.2)
        st = self.d.styles["Normal"]; st.font.name = BODY; st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
        st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.18
        hp = s.header.paragraphs[0]; r = hp.add_run(pack.upper()); _font(r, HEAD, 7.5, True, color=ACC)
        hp.add_run("\t\t"); r = hp.add_run(titulo); _font(r, HEAD, 7.5, color=GRAY); _border(hp, "bottom", "D8DCE3", 4)
        fp = s.footer.paragraphs[0]; r = fp.add_run(AVISO + "\t"); _font(r, HEAD, 7.5, italic=True, color=GRAY)
        r = fp.add_run("Pág. "); _font(r, HEAD, 8, color=GRAY); _field(fp, "PAGE")
        p = self.d.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(titulo); _font(r, HEAD, 22, True, color=NAVY)
        if sub:
            p = self.d.add_paragraph(); r = p.add_run(sub); _font(r, HEAD, 10.5, color=GRAY); _border(p, "bottom", "1F2A44", 12)
            p.paragraph_format.space_after = Pt(12)
    def box(self, titulo, lineas, fill="EEF1F6"):
        t = self.d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]; _shade(c._tc, fill)
        p = c.paragraphs[0]; r = p.add_run(titulo.upper()); _font(r, HEAD, 8.5, True, color=ACC)
        for l in lineas:
            p = c.add_paragraph(); p.paragraph_format.space_after = Pt(2); rich(p, "•  " + l, 9.5, name=HEAD)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
    def h(self, text):
        p = self.d.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.paragraph_format.keep_with_next = True
        r = p.add_run(text.upper()); _font(r, HEAD, 10.5, True, color=NAVY); _border(p, "bottom", "D8DCE3", 4)
    def p(self, text, just=True, size=10.5, align=None, italic=None, color=None):
        p = self.d.add_paragraph()
        if just: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if align == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rich(p, text, size, color=color, italic=italic); return p
    def bullets(self, items):
        for i in items:
            p = self.d.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.first_line_indent = Cm(-0.4)
            p.paragraph_format.space_after = Pt(3); rich(p, "–  " + i)
    def clauses(self, items, start=0):
        for n, it in enumerate(items, start):
            tit, txt = it[0], it[1]
            p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r = p.add_run(f"{ORD[n]}.- {tit}. " if tit else f"{ORD[n]}.- "); _font(r, BODY, 10.5, True, color=NAVY)
            rich(p, txt)
            if len(it) > 2: self.sub(it[2])
    def sub(self, items):
        for i in items:
            p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.5); p.paragraph_format.space_after = Pt(3)
            rich(p, i)
    def table(self, header, rows, widths=None, size=9.5):
        t = self.d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]; _shade(c._tc, "1F2A44"); c.paragraphs[0].text = ""
            r = c.paragraphs[0].add_run(h); _font(r, HEAD, size, True, color=RGBColor(255,255,255))
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].paragraphs[0].paragraph_format.space_after = Pt(2); rich(cells[i].paragraphs[0], v, size, name=HEAD)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths): row.cells[i].width = Cm(w)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
    def firmas(self, partes):
        self.d.add_paragraph()
        t = self.d.add_table(rows=1, cols=len(partes))
        for i, (rol) in enumerate(partes):
            c = t.rows[0].cells[i]; p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run("\n\n______________________________"); _font(r, BODY, 10)
            q = c.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; rich(q, f"**{rol}**\nFirma y aclaración: [[nombre]]\nDNI/CUIT: [[número]]", 9.5, name=HEAD)
    def pagebreak(self): self.d.add_page_break()
    def save(self, path): self.d.save(path)
