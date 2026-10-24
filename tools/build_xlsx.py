import datetime as dt
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "private", "kit.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
AZUL = "1F4E79"; CLARO = "DDEBF7"; INPUT = "FFF2CC"
H = Font(bold=True, color="FFFFFF"); B = Font(bold=True)
fill = lambda c: PatternFill("solid", fgColor=c)
thin = Side(style="thin", color="BFBFBF"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
ARS = '"$" #,##0.00'; PCT = "0.0%"

CATS = [("A",12009410.45,49527.18,49527.18),("B",17595182.74,56379.08,56379.08),
("C",24670494.31,66020.12,64530.58),("D",30628651.43,84612.93,82564.81),
("E",36028231.33,119811.45,108267.51),("F",45151659.41,150784.21,129930.65),
("G",53995798.87,230312.94,158815.05),("H",81924660.37,522706.68,317895.01),
("I",91699761.90,963747.86,474992.78),("J",105012519.20,1167299.76,580793.69),
("K",126610838.75,1614446.04,702103.24)]

wb = Workbook()

def header(ws, row, titles, widths=None):
    for i, t in enumerate(titles, 1):
        c = ws.cell(row=row, column=i, value=t); c.font = H; c.fill = fill(AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = box
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w

def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = Font(bold=True, size=16, color=AZUL)
    if sub: ws["A2"] = sub; ws["A2"].font = Font(italic=True, color="595959")

ws = wb.active; ws.title = "Inicio"
title(ws, "Kit Monotributista 2026", "Control de facturación, categoría y finanzas personales. Funciona en Excel y Google Sheets.")
ws.column_dimensions["A"].width = 110
lines = [
 "CÓMO USARLO (5 minutos)",
 "1. Pestaña 'Facturación': cargá cada factura emitida (fecha e importe). Es lo único que tenés que cargar seguido.",
 "2. Pestaña 'Control': elegí el mes de cierre (celda amarilla). Ves al instante tu facturación de los últimos 12 meses,",
 "   la categoría que te corresponde, cuánto margen te queda antes de saltar de categoría y si estás en riesgo de exclusión.",
 "3. Pestaña 'Finanzas': cargá tus gastos por mes y el tipo de cambio. Ves ahorro en pesos y en dólares, y cuánto pesa la cuota.",
 "4. Pestaña 'Parámetros': escalas vigentes de agosto 2026 a enero 2027. Cuando ARCA publique nuevas escalas, actualizá solo esa tabla.",
 "",
 "Celdas amarillas = las cargás vos. Todo lo demás se calcula solo. No borres filas de las tablas, solo su contenido.",
 "",
 "RECATEGORIZACIÓN",
 "Es semestral: en enero (con los 12 meses anteriores) y en julio. La próxima vence en los primeros días de febrero 2027",
 "y se calcula con la facturación de enero a diciembre 2026. Usá 'Control' con mes de cierre diciembre 2026 para saber qué categoría declarar.",
 "",
 "AVISO",
 "Esta planilla es una herramienta de organización personal. No reemplaza el asesoramiento de un contador ni la información oficial de ARCA",
 "(arca.gob.ar). Verificá siempre los valores vigentes antes de recategorizarte o pagar. Escalas tomadas de publicaciones de agosto 2026.",
 "Para uso personal. Prohibida su reventa o redistribución.",
]
for i, l in enumerate(lines, 4):
    ws.cell(row=i, column=1, value=l)
    if l.isupper() or l.startswith("CÓMO"): ws.cell(row=i, column=1).font = Font(bold=True, color=AZUL, size=12)

p = wb.create_sheet("Parámetros")
title(p, "Escalas del Monotributo", "Vigentes agosto 2026 – enero 2027. Editá esta tabla cuando cambien las escalas.")
header(p, 4, ["Categoría", "Tope ingresos brutos anuales", "Cuota mensual servicios", "Cuota mensual venta de bienes"], [12, 28, 24, 28])
for r, (c, tope, s, v) in enumerate(CATS, 5):
    p.cell(row=r, column=1, value=c).alignment = Alignment(horizontal="center")
    for col, val in ((2, tope), (3, s), (4, v)):
        cell = p.cell(row=r, column=col, value=val); cell.number_format = ARS; cell.fill = fill(INPUT)
    for col in range(1, 5): p.cell(row=r, column=col).border = box
p["A17"] = "Tipo de actividad:"; p["A17"].font = B
p["B17"] = "Servicios"; p["B17"].fill = fill(INPUT)
dv = DataValidation(type="list", formula1='"Servicios,Venta de bienes"', allow_blank=False); p.add_data_validation(dv); dv.add("B17")
for n, ref in (("Cats", "$A$5:$A$15"), ("Topes", "$B$5:$B$15"), ("CuotaServ", "$C$5:$C$15"), ("CuotaVenta", "$D$5:$D$15"), ("Actividad", "$B$17")):
    wb.defined_names[n] = DefinedName(n, attr_text=f"'Parámetros'!{ref}")

f = wb.create_sheet("Facturación")
title(f, "Registro de facturación", "Una fila por factura. Solo fecha e importe son obligatorios.")
header(f, 4, ["Fecha", "Nº comprobante", "Cliente", "Concepto", "Importe ($)"], [14, 18, 28, 36, 18])
N = 1000
for r in range(5, 5 + N):
    f.cell(row=r, column=1).number_format = "dd/mm/yyyy"
    f.cell(row=r, column=5).number_format = ARS
    for col in range(1, 6): f.cell(row=r, column=col).fill = fill(INPUT)
f["A5"] = dt.date(2026, 1, 15); f["B5"] = "C 0001-00000001"; f["C5"] = "Cliente ejemplo"; f["D5"] = "Borrá estas filas de ejemplo"; f["E5"] = 850000
f["A6"] = dt.date(2026, 2, 10); f["B6"] = "C 0001-00000002"; f["C6"] = "Cliente ejemplo"; f["D6"] = "Borrá estas filas de ejemplo"; f["E6"] = 920000
f.freeze_panes = "A5"
wb.defined_names["FFecha"] = DefinedName("FFecha", attr_text=f"'Facturación'!$A$5:$A${4+N}")
wb.defined_names["FImporte"] = DefinedName("FImporte", attr_text=f"'Facturación'!$E$5:$E${4+N}")

c = wb.create_sheet("Control")
title(c, "Control de categoría", "Elegí el mes de cierre (primer día del mes). Se analizan los 12 meses que terminan ahí.")
c.column_dimensions["A"].width = 38; c.column_dimensions["B"].width = 26; c.column_dimensions["C"].width = 4
c.column_dimensions["D"].width = 14; c.column_dimensions["E"].width = 20; c.column_dimensions["F"].width = 20
c["A4"] = "Mes de cierre"; c["B4"] = dt.date(2026, 12, 1); c["B4"].number_format = "mmm-yyyy"; c["B4"].fill = fill(INPUT)
rows = [
 ("Facturación últimos 12 meses", "=SUM(E5:E16)", ARS),
 ("Categoría que te corresponde", '=IF(B5>INDEX(Topes,11),"EXCLUIDO",IF(B5<=INDEX(Topes,1),"A",INDEX(Cats,MATCH(B5-0.001,Topes,1)+1)))', None),
 ("Tope de esa categoría", '=IF(B6="EXCLUIDO",0,INDEX(Topes,MATCH(B6,Cats,0)))', ARS),
 ("Margen antes de subir de categoría", '=IF(B6="EXCLUIDO",0,B7-B5)', ARS),
 ("% del tope usado", '=IF(B7=0,1,B5/B7)', PCT),
 ("Margen antes de exclusión (cat. K)", "=INDEX(Topes,11)-B5", ARS),
 ("Cuota mensual estimada", '=IF(B6="EXCLUIDO",0,INDEX(IF(Actividad="Servicios",CuotaServ,CuotaVenta),MATCH(B6,Cats,0)))', ARS),
 ("Cuota anual estimada", "=B11*12", ARS),
 ("Presión fiscal efectiva (cuota/facturación)", '=IF(B5=0,0,B12/B5)', PCT),
 ("Promedio mensual (últimos 12 meses)", "=B5/12", ARS),
 ("Promedio últimos 3 meses", "=SUM(E14:E16)/3", ARS),
 ("Proyección anual al ritmo de los últimos 3 meses", "=B15*12", ARS),
 ("Categoría proyectada", '=IF(B16>INDEX(Topes,11),"EXCLUIDO",IF(B16<=INDEX(Topes,1),"A",INDEX(Cats,MATCH(B16-0.001,Topes,1)+1)))', None),
 ("Estado", '=IF(B6="EXCLUIDO","⚠ Superaste el tope máximo: consultá a un contador",IF(B9>=0.9,"⚠ Estás a menos del 10% del tope de tu categoría",IF(B17<>B6,"ℹ Al ritmo actual vas a cambiar a categoría "&B17,"✓ Todo en orden")))', None),
]
for i, (lab, fo, fmt) in enumerate(rows, 5):
    c.cell(row=i, column=1, value=lab).font = B
    cell = c.cell(row=i, column=2, value=fo); cell.fill = fill(CLARO); cell.border = box
    if fmt: cell.number_format = fmt
c["B18"].alignment = Alignment(wrap_text=True)
c.merge_cells("B18:F18"); c.row_dimensions[18].height = 30
for col, t in zip("DEF", ["Mes", "Facturado", "Acumulado"]):
    cell = c[f"{col}4"]; cell.value = t; cell.font = H; cell.fill = fill(AZUL); cell.alignment = Alignment(horizontal="center")
for i in range(12):
    r = 5 + i
    c[f"D{r}"] = f"=EDATE($B$4,{i-11})"; c[f"D{r}"].number_format = "mmm-yy"
    c[f"E{r}"] = f'=SUMIFS(FImporte,FFecha,">="&D{r},FFecha,"<="&EOMONTH(D{r},0))'; c[f"E{r}"].number_format = ARS
    c[f"F{r}"] = f"=SUM($E$5:E{r})"; c[f"F{r}"].number_format = ARS
    for col in "DEF": c[f"{col}{r}"].border = box
c.conditional_formatting.add("B9", CellIsRule(operator="greaterThanOrEqual", formula=["0.9"], fill=fill("F8CBAD")))
c.conditional_formatting.add("B18", FormulaRule(formula=['LEFT($B$18,1)="⚠"'], fill=fill("F8CBAD")))
c.conditional_formatting.add("B18", FormulaRule(formula=['LEFT($B$18,1)="✓"'], fill=fill("C6EFCE")))

g = wb.create_sheet("Finanzas")
title(g, "Finanzas personales del monotributista", "Ingresos vienen solos de 'Facturación'. Cargá gastos y tipo de cambio (celdas amarillas).")
g["A4"] = "Año"; g["B4"] = 2026; g["B4"].fill = fill(INPUT); g["A4"].font = B
meses = ["Ene","Feb","Mar","Abr","May","Jun","Jul","Ago","Sep","Oct","Nov","Dic"]
g.column_dimensions["A"].width = 34
for j, m in enumerate(meses):
    col = chr(66 + j)
    g.column_dimensions[col].width = 13
    cell = g[f"{col}6"]; cell.value = m; cell.font = H; cell.fill = fill(AZUL); cell.alignment = Alignment(horizontal="center")
g.column_dimensions["N"].width = 15; g["N6"] = "Total"; g["N6"].font = H; g["N6"].fill = fill(AZUL)
g["A6"] = "Concepto"; g["A6"].font = H; g["A6"].fill = fill(AZUL)
gastos = ["Alquiler / vivienda", "Servicios (luz, gas, internet)", "Supermercado", "Transporte", "Salud", "Educación", "Salidas y ocio", "Suscripciones", "Tarjeta de crédito", "Otros"]
labels = [("Ingresos facturados", "ing"), ("Otros ingresos", "in"), ("Cuota monotributo", "cuota"), ("Ingresos brutos (IIBB provincial)", "in")] + [(x, "in") for x in gastos] + \
         [("Total gastos", "tg"), ("Ahorro del mes ($)", "ah"), ("Tasa de ahorro", "ta"), ("Tipo de cambio (ARS por USD)", "in"), ("Ahorro del mes (USD)", "usd")]
rowmap = {}
for i, (lab, kind) in enumerate(labels, 7):
    rowmap[lab] = i; g.cell(row=i, column=1, value=lab).font = B if kind != "in" else Font()
r_ing, r_otros, r_cuota, r_tg = rowmap["Ingresos facturados"], rowmap["Otros ingresos"], rowmap["Cuota monotributo"], rowmap["Total gastos"]
r_ah, r_ta, r_tc, r_usd = rowmap["Ahorro del mes ($)"], rowmap["Tasa de ahorro"], rowmap["Tipo de cambio (ARS por USD)"], rowmap["Ahorro del mes (USD)"]
for j in range(12):
    col = chr(66 + j); d = f"DATE($B$4,{j+1},1)"
    for lab, kind in labels:
        r = rowmap[lab]; cell = g[f"{col}{r}"]; cell.border = box
        if kind == "ing": cell.value = f'=SUMIFS(FImporte,FFecha,">="&{d},FFecha,"<="&EOMONTH({d},0))'
        elif kind == "cuota": cell.value = '=Control!$B$11'; cell.fill = fill(INPUT)
        elif kind == "in": cell.fill = fill(INPUT)
        elif kind == "tg": cell.value = f"=SUM({col}{r_cuota}:{col}{r_tg-1})"
        elif kind == "ah": cell.value = f"={col}{r_ing}+{col}{r_otros}-{col}{r_tg}"
        elif kind == "ta": cell.value = f"=IF(({col}{r_ing}+{col}{r_otros})=0,0,{col}{r_ah}/({col}{r_ing}+{col}{r_otros}))"
        elif kind == "usd": cell.value = f"=IF({col}{r_tc}>0,{col}{r_ah}/{col}{r_tc},0)"
        cell.number_format = PCT if kind == "ta" else ('"US$" #,##0.00' if kind == "usd" else ARS)
for lab, kind in labels:
    r = rowmap[lab]; cell = g[f"N{r}"]; cell.border = box; cell.font = B
    if kind == "ta": cell.value = f"=IF((N{r_ing}+N{r_otros})=0,0,N{r_ah}/(N{r_ing}+N{r_otros}))"; cell.number_format = PCT
    elif lab.startswith("Tipo de cambio"): cell.value = f'=IFERROR(AVERAGEIF(B{r}:M{r},">0"),0)'; cell.number_format = ARS
    else: cell.value = f"=SUM(B{r}:M{r})"; cell.number_format = '"US$" #,##0.00' if kind == "usd" else ARS
g.cell(row=r_cuota, column=1).value = "Cuota monotributo (editable)"
g.freeze_panes = "B7"
for ws in wb.worksheets: ws.sheet_view.showGridLines = False
wb.active = 0
wb.save(OUT); print("ok", OUT)
