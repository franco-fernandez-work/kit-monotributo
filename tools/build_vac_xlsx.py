import os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "private", "vacaciones.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
AZUL = "1F4E79"; CLARO = "DDEBF7"; INPUT = "FFF2CC"; ROJO = "F8CBAD"; VERDE = "C6E0B4"
H = Font(bold=True, color="FFFFFF"); B = Font(bold=True)
fill = lambda c: PatternFill("solid", fgColor=c)
thin = Side(style="thin", color="BFBFBF"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
ARS = '"$" #,##0.00'; FECHA = "dd/mm/yyyy"; MES = "mmm yy"
N0, N1 = 7, 56          # filas de empleados en Equipo (50 personas)
D = datetime.date

wb = Workbook()

def title(ws, text, sub):
    ws["A1"] = text; ws["A1"].font = Font(bold=True, size=16, color=AZUL)
    ws["A2"] = sub; ws["A2"].font = Font(italic=True, color="595959")

def cab(ws, row, heads, widths, col0=1):
    for i, (t, w) in enumerate(zip(heads, widths), col0):
        c = ws.cell(row=row, column=i, value=t); c.font = H; c.fill = fill(AZUL); c.border = box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[ws.cell(row=row, column=i).column_letter].width = w
    ws.row_dimensions[row].height = 44

# ---------- Inicio ----------
ws = wb.active; ws.title = "Inicio"
title(ws, "Planilla de Vacaciones 2026-2027", "Cronograma anual del equipo, días y monto de cada persona y aviso de licencia. Excel y Google Sheets.")
ws.column_dimensions["A"].width = 118
lines = [
 "CÓMO USARLO",
 "1. Pestaña 'Equipo': cargá nombre, fecha de ingreso y sueldo mensual bruto. Te dice cuántos días le tocan, el valor del día y el monto.",
 "2. En la misma fila poné cuándo se va de vacaciones (hasta dos tramos). Te avisa si se pasa de sus días o si cae fuera de temporada.",
 "3. Pestaña 'Cronograma': de octubre a septiembre, cuántos días de vacaciones tiene cada persona por mes y cuántas están afuera a la vez.",
 "4. Pestaña 'Aviso': elegí una persona y te arma la notificación para imprimir, con la fecha límite para entregarla.",
 "5. Cambiá el año en 'Equipo' (celda B3) cuando empiece el período nuevo. Todo se recalcula.",
 "",
 "Celdas amarillas = las cargás vos. Celdas celestes = se calculan solas. Las 4 filas de ejemplo se pisan o se borran.",
 "",
 "REGLAS QUE APLICA (Ley de Contrato de Trabajo 20.744 y reforma Ley 27.802)",
 "Días (art. 150): 14 corridos hasta 5 años de antigüedad, 21 hasta 10, 28 hasta 20 y 35 con más de 20. La antigüedad se mide al 31 de diciembre.",
 "Menos de 6 meses en el año (art. 151): un día de descanso por cada 20 días trabajados.",
 "Monto (art. 155): sueldo mensual / 25 por cada día de vacaciones, y se paga antes de que empiece la licencia.",
 "Temporada y aviso (art. 154, texto Ley 27.802): se otorgan entre el 1 de octubre y el 30 de abril, salvo acuerdo; el aviso es por escrito, con 30 días de anticipación como mínimo.",
 "Fraccionamiento: solo por acuerdo con la persona, en tramos de al menos 7 días corridos. Los días son corridos: cuentan sábados, domingos y feriados.",
 "",
 "AVISO",
 "Herramienta de cálculo orientativa para sueldos mensuales. Los convenios colectivos, el sueldo con comisiones o a jornal y los cambios",
 "de ley pueden modificar los resultados. Ante dudas consultá a un contador o abogado laboral. Para uso propio, prohibida su reventa.",
]
for i, l in enumerate(lines, 4):
    ws.cell(row=i, column=1, value=l)
    if l.startswith(("CÓMO", "REGLAS", "AVISO")): ws.cell(row=i, column=1).font = Font(bold=True, color=AZUL, size=12)

# ---------- Equipo ----------
e = wb.create_sheet("Equipo")
title(e, "Equipo y vacaciones", "Una fila por persona. Amarillo = lo cargás vos.")
e["A3"] = "Año del período"; e["A3"].font = B
e["B3"] = 2026; e["B3"].fill = fill(INPUT); e["B3"].border = box
e["C3"] = "Las vacaciones del año X se toman entre el 1/10 del año X y el 30/4 del año siguiente."
e["C3"].font = Font(color="595959")
e["A4"] = "Aviso con anticipación (días)"; e["A4"].font = B
e["B4"] = 30; e["B4"].fill = fill(INPUT); e["B4"].border = box
wb.defined_names["Anio"] = DefinedName("Anio", attr_text="'Equipo'!$B$3")
wb.defined_names["AvisoDias"] = DefinedName("AvisoDias", attr_text="'Equipo'!$B$4")
cab(e, 6, ["Nombre", "Fecha de ingreso", "Sueldo mensual bruto", "Años de antigüedad al 31/12", "Días que le corresponden",
           "Valor del día (sueldo / 25)", "Monto total de vacaciones", "Tramo 1 desde", "Tramo 1 hasta", "Tramo 2 desde",
           "Tramo 2 hasta", "Días programados", "Saldo sin programar", "Avisar antes del", "Estado"],
    [28, 14, 18, 14, 14, 17, 19, 13, 13, 13, 13, 12, 12, 14, 34])
corte = "DATE(Anio,12,31)"
trab = f"({corte}-MAX($B{{r}},DATE(Anio,1,1))+1)"
ejemplo = [
    ("Ejemplo: Lucía Fernández", D(2012, 2, 6), 1850000, D(2026, 12, 14), D(2027, 1, 3), None, None),
    ("Ejemplo: Martín Gómez", D(2021, 8, 17), 1420000, D(2027, 1, 11), D(2027, 1, 24), D(2027, 2, 8), D(2027, 2, 14)),
    ("Ejemplo: Carla Ruiz", D(2024, 3, 4), 1260000, D(2027, 1, 18), D(2027, 1, 31), None, None),
    ("Ejemplo: Tomás Díaz", D(2026, 9, 1), 1100000, D(2026, 12, 21), D(2026, 12, 26), None, None),
]
for r in range(N0, N1 + 1):
    t = trab.format(r=r)
    f = {
     4: f'=IF(OR($A{r}="",$B{r}=""),"",IFERROR(DATEDIF($B{r},{corte},"y"),0))',
     5: f'=IF(OR($A{r}="",$B{r}=""),"",IFERROR(IF({t}<183,INT({t}/20),IF({corte}<=EDATE($B{r},60),14,IF({corte}<=EDATE($B{r},120),21,IF({corte}<=EDATE($B{r},240),28,35)))),""))',
     6: f'=IF(OR($C{r}="",$E{r}=""),"",$C{r}/25)',
     7: f'=IF($F{r}="","",$F{r}*$E{r})',
     12: f'=IF($A{r}="","",IF(AND(N($H{r})>0,N($I{r})>=N($H{r})),$I{r}-$H{r}+1,0)+IF(AND(N($J{r})>0,N($K{r})>=N($J{r})),$K{r}-$J{r}+1,0))',
     13: f'=IF(OR($E{r}="",$L{r}=""),"",$E{r}-$L{r})',
     14: f'=IF(N($H{r})>0,$H{r}-AvisoDias,"")',
     15: (f'=IF($A{r}="","",IF($E{r}="","Falta la fecha de ingreso",IF($L{r}=0,"Sin programar",IF($M{r}<0,"Se pasa por "&-$M{r}&" días",'
          f'IF(OR(MIN(IF(N($H{r})>0,$H{r},99999),IF(N($J{r})>0,$J{r},99999))<DATE(Anio,10,1),MAX(N($I{r}),N($K{r}))>DATE(Anio+1,4,30)),"Fuera de temporada: requiere acuerdo",'
          f'IF($M{r}>0,"Programado, faltan "&$M{r}&" días","Completo"))))))'),
    }
    for col in range(1, 16):
        c = e.cell(row=r, column=col); c.border = box
        if col in f: c.value = f[col]; c.fill = fill(CLARO)
        else: c.fill = fill(INPUT)
    for col in (2, 8, 9, 10, 11, 14): e.cell(row=r, column=col).number_format = FECHA
    for col in (3, 6, 7): e.cell(row=r, column=col).number_format = ARS
for i, (n, ing, s, a, b, c2, d2) in enumerate(ejemplo, N0):
    for col, v in zip((1, 2, 3, 8, 9, 10, 11), (n, ing, s, a, b, c2, d2)): e.cell(row=i, column=col, value=v)
e.freeze_panes = "B7"
rojo = DifferentialStyle(fill=PatternFill(bgColor=ROJO, fill_type="solid"), font=Font(color="9C0006"))
verde = DifferentialStyle(fill=PatternFill(bgColor=VERDE, fill_type="solid"))
e.conditional_formatting.add(f"O{N0}:O{N1}", Rule(type="expression", dxf=rojo, formula=[f'OR(LEFT($O{N0},5)="Se pa",LEFT($O{N0},5)="Fuera")']))
e.conditional_formatting.add(f"O{N0}:O{N1}", Rule(type="expression", dxf=verde, formula=[f'$O{N0}="Completo"']))
e.conditional_formatting.add(f"N{N0}:N{N1}", Rule(type="expression", dxf=rojo, formula=[f'AND(N($N{N0})>0,TODAY()>$N{N0},TODAY()<N($H{N0}))']))
dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True, showErrorMessage=True,
                    errorTitle="Fecha", error="Escribí una fecha, por ejemplo 14/12/2026.")
e.add_data_validation(dv)
for rng in (f"B{N0}:B{N1}", f"H{N0}:K{N1}"): dv.add(rng)
e["C5"] = "Rojo en 'Avisar antes del' = la fecha límite pasó y las vacaciones todavía no empezaron."; e["C5"].font = Font(color="595959")

# ---------- Cronograma ----------
c = wb.create_sheet("Cronograma")
title(c, "Cronograma de vacaciones", "Días de vacaciones por persona y por mes, de octubre a septiembre. Se llena solo desde 'Equipo'.")
c["A3"] = "Año del período"; c["A3"].font = B; c["B3"] = "=Anio"
heads = ["Nombre"] + [None] * 12 + ["Total días"]
cab(c, 5, ["Nombre"] + [""] * 12 + ["Total días"], [28] + [9] * 12 + [11])
for j in range(12):
    cell = c.cell(row=5, column=2 + j, value=f"=EDATE(DATE(Anio,10,1),{j})"); cell.number_format = MES
for r in range(N0, N1 + 1):
    row = r - N0 + 6
    c.cell(row=row, column=1, value=f'=IF(Equipo!$A{r}="","",Equipo!$A{r})').border = box
    for j in range(12):
        col = c.cell(row=5, column=2 + j).column_letter
        m = f"{col}$5"
        v = (f'=IF(Equipo!$A{r}="","",MAX(0,MIN(N(Equipo!$I{r}),EOMONTH({m},0))-MAX(N(Equipo!$H{r}),{m})+1)'
             f'+MAX(0,MIN(N(Equipo!$K{r}),EOMONTH({m},0))-MAX(N(Equipo!$J{r}),{m})+1))')
        cc = c.cell(row=row, column=2 + j, value=v); cc.border = box; cc.alignment = Alignment(horizontal="center")
    t = c.cell(row=row, column=14, value=f'=IF(A{row}="","",SUM(B{row}:M{row}))'); t.border = box; t.font = B; t.alignment = Alignment(horizontal="center")
last = N1 - N0 + 6
fr = last + 2
c.cell(row=fr, column=1, value="Personas con vacaciones en el mes").font = B
c.cell(row=fr + 1, column=1, value="Días de vacaciones del equipo").font = B
for j in range(12):
    col = c.cell(row=5, column=2 + j).column_letter
    a = c.cell(row=fr, column=2 + j, value=f'=COUNTIF({col}6:{col}{last},">0")'); a.alignment = Alignment(horizontal="center"); a.fill = fill(CLARO)
    b = c.cell(row=fr + 1, column=2 + j, value=f"=SUM({col}6:{col}{last})"); b.alignment = Alignment(horizontal="center"); b.fill = fill(CLARO)
c.conditional_formatting.add(f"B6:M{last}", Rule(type="expression", dxf=DifferentialStyle(fill=PatternFill(bgColor="9DC3E6", fill_type="solid"), font=Font(bold=True)), formula=['AND(ISNUMBER(B6),B6>0)']))
c.conditional_formatting.add(f"B{fr}:M{fr}", Rule(type="expression", dxf=rojo, formula=[f"B{fr}>=3"]))
c["D3"] = "Rojo abajo = 3 o más personas de vacaciones a la vez en ese mes."; c["D3"].font = Font(color="595959")
c.freeze_panes = "B6"

# ---------- Aviso ----------
a = wb.create_sheet("Aviso")
title(a, "Notificación de vacaciones", "Elegí la persona. Imprimí la hoja, firmala y pedí la firma de recibido.")
a.column_dimensions["A"].width = 52; a.column_dimensions["B"].width = 30; a.column_dimensions["C"].width = 50
a["A4"] = "Persona"; a["A4"].font = B
a["B4"] = f"=Equipo!A{N0}"; a["B4"].fill = fill(INPUT); a["B4"].border = box
dvl = DataValidation(type="list", formula1=f"=Equipo!$A${N0}:$A${N1}", allow_blank=False); a.add_data_validation(dvl); dvl.add("B4")
a["A5"] = "Fecha de esta notificación"; a["A5"].font = B
a["B5"] = "=TODAY()"; a["B5"].number_format = FECHA; a["B5"].fill = fill(INPUT); a["B5"].border = box
a["C5"] = "Podés escribir otra fecha."; a["C5"].font = Font(color="595959")
a["A6"] = "Fila en 'Equipo'"; a["B6"] = f'=IFERROR(MATCH(B4,Equipo!$A${N0}:$A${N1},0),"")'; a["A6"].font = Font(color="808080"); a["B6"].font = Font(color="808080")
def idx(col): return f'IFERROR(INDEX(Equipo!${col}${N0}:${col}${N1},$B$6),"")'
rows = [
 ("Sr./Sra.", f'=B4', None),
 ("Le comunicamos que gozará de su licencia anual ordinaria (art. 154 LCT, texto Ley 27.802) correspondiente al año", "=Anio", "0"),
 ("Días corridos que le corresponden", "=" + idx("E"), "0"),
 ("Tramo 1: desde el", f'=IF(N({idx("H")})>0,{idx("H")},"")', FECHA),
 ("Tramo 1: hasta el", f'=IF(N({idx("I")})>0,{idx("I")},"")', FECHA),
 ("Se reintegra el", f'=IF(N({idx("I")})>0,{idx("I")}+1,"")', FECHA),
 ("Tramo 2: desde el", f'=IF(N({idx("J")})>0,{idx("J")},"")', FECHA),
 ("Tramo 2: hasta el", f'=IF(N({idx("K")})>0,{idx("K")},"")', FECHA),
 ("Monto a pagar antes de que empiece (art. 155)", f'=IF(OR({idx("F")}="",N({idx("H")})=0),"",{idx("F")}*(N({idx("I")})-N({idx("H")})+1))', ARS),
 ("Fecha límite para entregar este aviso", f'=IF(N({idx("H")})>0,{idx("H")}-AvisoDias,"")', FECHA),
 ("Control de anticipación", f'=IF(OR(N({idx("H")})=0,B5=""),"",IF({idx("H")}-B5>=AvisoDias,"Cumple: faltan "&({idx("H")}-B5)&" días","No cumple: faltan "&({idx("H")}-B5)&" días, hacen falta "&AvisoDias))', None),
]
for i, (lab, f, fmt) in enumerate(rows, 8):
    a.cell(row=i, column=1, value=lab).font = B
    cc = a.cell(row=i, column=2, value=f); cc.border = box; cc.fill = fill(CLARO); cc.alignment = Alignment(horizontal="left")
    if fmt: cc.number_format = fmt
a.conditional_formatting.add("B18", Rule(type="expression", dxf=rojo, formula=['LEFT($B$18,2)="No"']))
a["A21"] = "Firma del empleador: ______________________________"; a["A23"] = "Recibí la presente. Firma y fecha: ______________________________"
a["A25"] = "Entregá el original a la persona y guardá una copia firmada: sirve de prueba de que avisaste a tiempo."
a["A25"].font = Font(italic=True, color="595959")
for ws_ in (e, c, a):
    ws_.sheet_view.showGridLines = False
a.page_setup.orientation = "portrait"; a.page_setup.fitToWidth = 1
wb.save(OUT)
print("ok", OUT)
