"""Planilla 'Stock, costos y ventas para la temporada' (HUECOS.md #28, ronda 6). `python tools/build_stock_xlsx.py` arma private/stock.xlsx (solo local, necesita openpyxl; commitear el .xlsx)."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

priv = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "private")
AZUL = "1F4E79"; INP = "FFF2CC"; CAL = "DDEBF7"
H = Font(bold=True, color="FFFFFF"); fill = lambda c: PatternFill("solid", fgColor=c)
thin = Side(style="thin", color="BFBFBF"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
ARS = '"$" #,##0'; NP = 40; NV = 500; P0 = 5; V0 = 5   # filas de productos y de ventas

def titulo(ws, t, s):
    ws["A1"] = t; ws["A1"].font = Font(bold=True, size=16, color=AZUL); ws["A2"] = s; ws["A2"].font = Font(italic=True, color="595959")
def cab(ws, row, heads, widths):
    for i, (t, w) in enumerate(zip(heads, widths), 1):
        c = ws.cell(row=row, column=i, value=t); c.font = H; c.fill = fill(AZUL); c.border = box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 32
def st(c, kind, fmt=None):
    c.fill = fill(INP if kind == "i" else CAL); c.border = box
    if fmt: c.number_format = fmt

wb = Workbook(); ws = wb.active; ws.title = "Inicio"
titulo(ws, "Stock, costos y ventas para la temporada", "Para saber cuánto ganás de verdad, qué reponer y cuánta plata tenés parada en mercadería. Excel y Google Sheets.")
ws.column_dimensions["A"].width = 120
lines = ["CÓMO USARLA",
 "1. Pestaña 'Parámetros': cargá tus costos fijos del período (puesto, internet, cuota de monotributo, packaging) y tu meta de ganancia.",
 "2. Pestaña 'Productos': una fila por producto con lo que te cuesta, a cuánto lo vendés, el % de comisión y envío, el stock que compraste y tu stock mínimo. El margen se calcula solo.",
 "3. Pestaña 'Ventas': una fila por venta (fecha, producto, cantidad). Elegí el producto de la lista. Si vendiste a otro precio, escribilo en 'Precio real'.",
 "4. De vuelta en 'Productos' ves cuántas unidades vendiste, cuánto stock te queda y el aviso REPONER cuando llegás al mínimo.",
 "5. Pestaña 'Resumen': ingresos, costo de la mercadería vendida, comisiones, ganancia neta, plata parada en stock y producto más rentable.",
 "", "Celdas amarillas = las cargás vos. Celdas celestes = se calculan solas. Los productos de ejemplo son referencia: pisalos con los tuyos (hasta 40 productos y 500 ventas).",
 "El % de comisión y envío lo calculás vos como porcentaje del precio (por ejemplo: la comisión de la plataforma de cobro o de la tienda online, más lo que absorbés de envío).",
 "", "AVISO", "Herramienta de cálculo orientativa. No es asesoramiento contable ni impositivo. No incluye impuestos: consultá con tu contador. Para uso propio, prohibida su reventa."]
for i, l in enumerate(lines, 4):
    ws.cell(row=i, column=1, value=l)
    if l in ("CÓMO USARLA", "AVISO"): ws.cell(row=i, column=1).font = Font(bold=True, color=AZUL, size=12)

wp = wb.create_sheet("Parámetros"); titulo(wp, "Parámetros", "Costos fijos del período y meta")
wp.column_dimensions["A"].width = 46; wp.column_dimensions["B"].width = 18
filas = [("Costos fijos del período (puesto, internet, monotributo, packaging)", 150000, ARS), ("Meta de ganancia neta del período", 400000, ARS)]
for i, (n, v, f) in enumerate(filas, 4):
    wp.cell(row=i, column=1, value=n).font = Font(bold=True); st(wp.cell(row=i, column=2, value=v), "i", f)
wp["A7"] = "Los costos fijos se pagan venda lo que venda. El Resumen los descuenta de la ganancia bruta."; wp["A7"].font = Font(italic=True, color="595959")

wq = wb.create_sheet("Productos"); titulo(wq, "Productos y stock", "Costo, precio, stock inicial y mínimo; el resto se calcula")
cab(wq, 4, ["Producto", "Costo unitario", "Precio de venta", "Comisión y envío (% del precio)", "Stock comprado", "Stock mínimo (aviso)", "Margen $ por unidad", "Margen %", "Unidades vendidas", "Stock actual", "Estado", "Ganancia del producto"],
    [30, 14, 14, 16, 12, 12, 14, 10, 12, 11, 12, 16])
ej = [("Vela aromática", 3500, 8500, 0.05, 60, 10), ("Caja de alfajores x6", 4200, 9000, 0.05, 80, 15), ("Cuadro decorativo", 6000, 15000, 0.08, 30, 5), ("Pulsera artesanal", 1500, 4500, 0.05, 100, 20), ("Set de cuadernos", 5200, 12000, 0.05, 40, 8)]
for k in range(NP):
    r = P0 + k
    for col in range(1, 7):
        c = wq.cell(row=r, column=col)
        if k < len(ej): c.value = ej[k][col - 1]
        st(c, "i", {2: ARS, 3: ARS, 4: "0.0%"}.get(col))
    rv = f"Ventas!$B${V0}:$B${V0 + NV - 1}"; rc = f"Ventas!$C${V0}:$C${V0 + NV - 1}"
    st(wq.cell(row=r, column=7, value=f'=IF(A{r}="","",C{r}-B{r}-C{r}*D{r})'), "c", ARS)
    st(wq.cell(row=r, column=8, value=f'=IF(OR(A{r}="",C{r}=0),"",G{r}/C{r})'), "c", "0%")
    st(wq.cell(row=r, column=9, value=f'=IF(A{r}="","",SUMIF({rv},A{r},{rc}))'), "c", "0")
    st(wq.cell(row=r, column=10, value=f'=IF(A{r}="","",E{r}-I{r})'), "c", "0")
    st(wq.cell(row=r, column=11, value=f'=IF(A{r}="","",IF(J{r}<=0,"SIN STOCK",IF(J{r}<=F{r},"REPONER","OK")))'), "c")
    st(wq.cell(row=r, column=12, value=f'=IF(A{r}="","",G{r}*I{r})'), "c", ARS)
    wq.cell(row=r, column=11).alignment = Alignment(horizontal="center")
last = P0 + NP - 1
wq.conditional_formatting.add(f"K{P0}:K{last}", CellIsRule(operator="equal", formula=['"REPONER"'], fill=fill("FFD966"), font=Font(bold=True)))
wq.conditional_formatting.add(f"K{P0}:K{last}", CellIsRule(operator="equal", formula=['"SIN STOCK"'], fill=fill("F4B084"), font=Font(bold=True)))
wq.freeze_panes = "B5"

wv = wb.create_sheet("Ventas"); titulo(wv, "Ventas", "Una fila por venta. Elegí el producto de la lista")
cab(wv, 4, ["Fecha", "Producto", "Cantidad", "Precio real (opcional)", "Importe"], [14, 30, 11, 18, 16])
dv = DataValidation(type="list", formula1=f"=Productos!$A${P0}:$A${last}", allow_blank=True); wv.add_data_validation(dv)
ejv = [("2026-12-01", "Vela aromática", 3), ("2026-12-02", "Caja de alfajores x6", 5), ("2026-12-03", "Pulsera artesanal", 8), ("2026-12-05", "Cuadro decorativo", 1)]
import datetime
for k in range(NV):
    r = V0 + k
    st(wv.cell(row=r, column=1), "i", "dd/mm/yyyy"); st(wv.cell(row=r, column=2), "i"); st(wv.cell(row=r, column=3), "i", "0"); st(wv.cell(row=r, column=4), "i", ARS)
    dv.add(wv.cell(row=r, column=2))
    if k < len(ejv):
        wv.cell(row=r, column=1, value=datetime.date.fromisoformat(ejv[k][0])); wv.cell(row=r, column=2, value=ejv[k][1]); wv.cell(row=r, column=3, value=ejv[k][2])
    st(wv.cell(row=r, column=5, value=f'=IF(B{r}="","",C{r}*IF(D{r}>0,D{r},IFERROR(VLOOKUP(B{r},Productos!$A${P0}:$C${last},3,FALSE),0)))'), "c", ARS)
wv.freeze_panes = "A5"

wz = wb.create_sheet("Resumen"); titulo(wz, "Resumen del período", "Cuánto ganaste y cuánta plata tenés parada")
wz.column_dimensions["A"].width = 52; wz.column_dimensions["B"].width = 22
pr = lambda col: f"Productos!${col}${P0}:${col}${last}"
filas = [("Ingresos por ventas", f"=SUM(Ventas!E{V0}:E{V0 + NV - 1})", ARS),
         ("Costo de la mercadería vendida", f"=SUMPRODUCT({pr('B')},{pr('I')})", ARS),
         ("Comisiones y envíos (estimado con el precio de lista)", f"=SUMPRODUCT({pr('C')},{pr('D')},{pr('I')})", ARS),
         ("Ganancia bruta", "=B5-B6-B7", ARS),
         ("Costos fijos del período", "=Parámetros!B4", ARS),
         ("GANANCIA NETA", "=B8-B9", ARS),
         ("Meta de ganancia", "=Parámetros!B5", ARS),
         ("Falta para la meta (0 = cumplida)", "=MAX(0,B11-B10)", ARS),
         ("Unidades vendidas", f"=SUM({pr('I')})", "0"),
         ("Plata invertida en mercadería comprada", f"=SUMPRODUCT({pr('B')},{pr('E')})", ARS),
         ("Plata parada en stock (a costo)", f"=SUMPRODUCT({pr('B')},{pr('J')})", ARS),
         ("Productos para reponer o sin stock", f'=COUNTIF({pr("K")},"REPONER")+COUNTIF({pr("K")},"SIN STOCK")', "0"),
         ("Producto más rentable", f'=IFERROR(INDEX({pr("A")},MATCH(MAX({pr("L")}),{pr("L")},0)),"")', None)]
for i, (n, fm, f) in enumerate(filas, 5):
    wz.cell(row=i, column=1, value=n).font = Font(bold=n.isupper() or n == "Ganancia bruta"); st(wz.cell(row=i, column=2, value=fm), "c", f)
wz["B10"].font = Font(bold=True, size=12)
wz["A19"] = "Si la ganancia neta es negativa, los costos fijos todavía no se pagaron: revisá precios, comisiones y stock parado."; wz["A19"].font = Font(italic=True, color="595959")
wb.save(os.path.join(priv, "stock.xlsx")); print("ok stock.xlsx")
