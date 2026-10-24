import os, json, sys, datetime, urllib.request
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule, FormulaRule

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "private", "alquileres.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
IPC_URL = os.environ.get("IPC_URL") or "https://apis.datos.gob.ar/series/api/series/?ids=148.3_INIVELNAL_DICI_M_26&format=json&limit=1000&sort=asc"
AZUL = "1F4E79"; CLARO = "DDEBF7"; INPUT = "FFF2CC"
H = Font(bold=True, color="FFFFFF"); B = Font(bold=True)
fill = lambda c: PatternFill("solid", fgColor=c)
thin = Side(style="thin", color="BFBFBF"); box = Border(top=thin, bottom=thin, left=thin, right=thin)
ARS = '"$" #,##0'; PCT = "0.0%"; MES = "mmm yyyy"; FECHA = "dd/mm/yyyy"

def ipc():
    # Serie oficial INDEC vía datos.gob.ar. Si no responde, la hoja queda para cargar a mano.
    try:
        # datos.gob.ar rechaza el User-Agent por defecto de urllib (403).
        req = urllib.request.Request(IPC_URL, headers={"User-Agent": "kit-monotributo-build/1.0"})
        with urllib.request.urlopen(req, timeout=30) as r:
            return [(datetime.date.fromisoformat(f), v) for f, v in json.load(r)["data"] if v is not None]
    except Exception as e:
        print("AVISO: no se pudo bajar el IPC:", e, file=sys.stderr)
        return []

wb = Workbook()

def title(ws, text, sub):
    ws["A1"] = text; ws["A1"].font = Font(bold=True, size=16, color=AZUL)
    ws["A2"] = sub; ws["A2"].font = Font(italic=True, color="595959")

def cabecera(ws, row, heads, widths):
    for i, (t, w) in enumerate(zip(heads, widths), 1):
        c = ws.cell(row=row, column=i, value=t); c.font = H; c.fill = fill(AZUL); c.border = box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[chr(64 + i)].width = w
    ws.row_dimensions[row].height = 32

# Inicio
ws = wb.active; ws.title = "Inicio"
title(ws, "Planilla de Alquileres 2026", "Contratos, aumentos por IPC y cobros en un solo lugar. Funciona en Excel y Google Sheets.")
ws.column_dimensions["A"].width = 115
lines = [
 "CÓMO USARLO",
 "1. Pestaña 'Contratos': cargá cada alquiler con fecha de inicio, monto inicial y cada cuántos meses se actualiza (celdas amarillas).",
 "   La planilla te muestra el alquiler vigente hoy, el aumento acumulado y cuándo es el próximo ajuste.",
 "2. Pestaña 'Cobros': anotá cada pago (alquiler, expensas, ABL, servicios) y marcá si está cobrado.",
 "3. Pestaña 'Resumen': lo cobrado y lo pendiente del año por propiedad.",
 "4. Pestaña 'IPC': el índice de precios del INDEC mes a mes. Viene cargado; cuando salga un mes nuevo,",
 "   agregá una fila con el primer día del mes y el valor del índice (indec.gob.ar, 'Índice de precios al consumidor').",
 "",
 "CÓMO CALCULA",
 "Alquiler vigente = monto inicial × IPC del mes base del último ajuste / IPC del mes base del inicio.",
 "Por defecto usa los IPC publicados hasta 2 meses antes de cada ajuste (rezago 1), que es lo habitual porque",
 "el INDEC publica cada mes a mediados del mes siguiente. Si tu contrato usa otro criterio, cambiá 'Rezago' en la pestaña IPC.",
 "",
 "AVISO",
 "Herramienta de cálculo personal. Lo que vale es lo que dice tu contrato: revisá su cláusula de actualización.",
 "Para uso personal. Prohibida su reventa o redistribución.",
]
for i, l in enumerate(lines, 4):
    ws.cell(row=i, column=1, value=l)
    if l in ("CÓMO USARLO", "CÓMO CALCULA", "AVISO"): ws.cell(row=i, column=1).font = Font(bold=True, color=AZUL, size=12)

# IPC
q = wb.create_sheet("IPC")
title(q, "IPC nivel general (INDEC)", "Índice base diciembre 2016 = 100. Agregá meses nuevos al final.")
q["D4"] = "Rezago (meses)"; q["D4"].font = B; q["E4"] = 1; q["E4"].fill = fill(INPUT); q["E4"].border = box
q["D5"] = "0 = usa el IPC hasta el mes anterior al ajuste. 1 = hasta dos meses antes."; q["D5"].font = Font(color="595959")
wb.defined_names["Rezago"] = DefinedName("Rezago", attr_text="'IPC'!$E$4")
cabecera(q, 4, ["Mes", "Índice", "Variación mensual"], [14, 14, 16])
datos = ipc()
FIN = 4 + 400
for r in range(5, FIN + 1):
    q[f"A{r}"].number_format = MES; q[f"B{r}"].number_format = "0.0000"; q[f"C{r}"].number_format = PCT
    q[f"C{r}"] = f'=IF(OR(B{r}="",B{r - 1}="",NOT(ISNUMBER(B{r - 1}))),"",B{r}/B{r - 1}-1)'
    if r - 5 >= len(datos):
        q[f"A{r}"].fill = fill(INPUT); q[f"B{r}"].fill = fill(INPUT)
for r, (f, v) in enumerate(datos, 5):
    q[f"A{r}"] = f; q[f"B{r}"] = v
q.freeze_panes = "A5"
for n, col in (("IPCMes", "A"), ("IPCIdx", "B")):
    wb.defined_names[n] = DefinedName(n, attr_text=f"'IPC'!${col}$5:${col}${FIN}")

# Contratos
c = wb.create_sheet("Contratos", 1)
title(c, "Contratos", "Cargá las celdas amarillas. El resto se calcula solo con la fecha de hoy.")
cabecera(c, 4, ["Propiedad", "Inquilino", "Inicio del contrato", "Monto inicial", "Ajuste cada (meses)", "Día de vencimiento",
                "Último ajuste", "Alquiler vigente", "Aumento acumulado", "Próximo ajuste", "Faltan (días)"],
         [26, 22, 14, 15, 12, 12, 14, 16, 14, 14, 11])
dv = DataValidation(type="list", formula1='"1,3,4,6,12"', allow_blank=True); c.add_data_validation(dv)
N = 50
for r in range(5, 5 + N):
    for col in "ABCDEF": c[f"{col}{r}"].fill = fill(INPUT)
    dv.add(f"E{r}")
    idx = lambda d: f"INDEX(IPCIdx,MATCH(EDATE({d},-(1+Rezago)),IPCMes,0))"
    c[f"G{r}"] = f'=IF(OR(C{r}="",E{r}=""),"",EDATE(C{r},E{r}*INT(DATEDIF(C{r},TODAY(),"m")/E{r})))'
    c[f"H{r}"] = f'=IF(G{r}="","",IF(G{r}=C{r},D{r},IFERROR(D{r}*{idx(f"G{r}")}/{idx(f"C{r}")},"Falta IPC")))'
    c[f"I{r}"] = f'=IF(ISNUMBER(H{r}),H{r}/D{r}-1,"")'
    c[f"J{r}"] = f'=IF(G{r}="","",EDATE(G{r},E{r}))'
    c[f"K{r}"] = f'=IF(J{r}="","",J{r}-TODAY())'
    for col, fmt in zip("CDGHIJK", [FECHA, ARS, FECHA, ARS, PCT, FECHA, "0"]):
        c[f"{col}{r}"].number_format = fmt
    for col in "ABCDEFGHIJK": c[f"{col}{r}"].border = box
c["A5"] = "Ejemplo: Depto Palermo"; c["B5"] = "Inquilino"; c["C5"] = datetime.date(2026, 1, 1); c["D5"] = 600000; c["E5"] = 3; c["F5"] = 10
c.freeze_panes = "B5"
c.conditional_formatting.add(f"K5:K{4 + N}", CellIsRule(operator="between", formula=["0", "15"], fill=fill("F8CBAD")))
c.conditional_formatting.add(f"H5:H{4 + N}", FormulaRule(formula=[f'H5="Falta IPC"'], fill=fill("F8CBAD")))
wb.defined_names["Props"] = DefinedName("Props", attr_text=f"'Contratos'!$A$5:$A${4 + N}")

# Cobros
p = wb.create_sheet("Cobros", 2)
title(p, "Cobros", "Una fila por cada pago que esperás. Marcá 'Sí' cuando lo cobres.")
cabecera(p, 4, ["Vence", "Propiedad", "Concepto", "Importe", "Cobrado", "Fecha de cobro", "Estado"], [13, 26, 16, 15, 10, 14, 14])
dvp = DataValidation(type="list", formula1="=Props", allow_blank=True)
dvc = DataValidation(type="list", formula1='"Alquiler,Expensas,ABL,Servicios,Otro"', allow_blank=True)
dvs = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
for d in (dvp, dvc, dvs): p.add_data_validation(d)
M = 1000
for r in range(5, 5 + M):
    for col in "ABCDEF": p[f"{col}{r}"].fill = fill(INPUT)
    dvp.add(f"B{r}"); dvc.add(f"C{r}"); dvs.add(f"E{r}")
    p[f"G{r}"] = f'=IF(A{r}="","",IF(E{r}="Sí","Cobrado",IF(A{r}<TODAY(),"Atrasado","Pendiente")))'
    for col, fmt in zip("ADF", [FECHA, ARS, FECHA]): p[f"{col}{r}"].number_format = fmt
    for col in "ABCDEFG": p[f"{col}{r}"].border = box
p["A5"] = datetime.date(2026, 11, 10); p["B5"] = "Ejemplo: Depto Palermo"; p["C5"] = "Alquiler"; p["D5"] = 600000; p["E5"] = "No"
p.freeze_panes = "A5"
p.conditional_formatting.add(f"G5:G{4 + M}", CellIsRule(operator="equal", formula=['"Atrasado"'], fill=fill("F8CBAD")))
p.conditional_formatting.add(f"G5:G{4 + M}", CellIsRule(operator="equal", formula=['"Cobrado"'], fill=fill("C6EFCE")))

# Resumen
s = wb.create_sheet("Resumen", 3)
title(s, "Resumen del año", "Cambiá el año en la celda amarilla.")
s["A4"] = "Año"; s["A4"].font = B; s["B4"] = 2026; s["B4"].fill = fill(INPUT); s["B4"].border = box
cabecera(s, 6, ["Propiedad", "Cobrado", "Pendiente", "Atrasado"], [26, 16, 16, 16])
rng = lambda col: f"Cobros!${col}$5:${col}${4 + M}"
anio = f'{rng("A")},">="&DATE($B$4,1,1),{rng("A")},"<"&DATE($B$4+1,1,1)'
for i, r in enumerate(range(7, 7 + N)):
    s[f"A{r}"] = f'=IF(Contratos!A{5 + i}="","",Contratos!A{5 + i})'
    for col, est in zip("BCD", ["Cobrado", "Pendiente", "Atrasado"]):
        s[f"{col}{r}"] = f'=IF($A{r}="","",SUMIFS({rng("D")},{rng("B")},$A{r},{rng("G")},"{est}",{anio}))'
        s[f"{col}{r}"].number_format = ARS
    for col in "ABCD": s[f"{col}{r}"].border = box
t = 7 + N
s[f"A{t}"] = "Total"; s[f"A{t}"].font = B
for col in "BCD":
    s[f"{col}{t}"] = f"=SUM({col}7:{col}{t - 1})"; s[f"{col}{t}"].number_format = ARS; s[f"{col}{t}"].font = B; s[f"{col}{t}"].fill = fill(CLARO)

for w in wb.worksheets: w.sheet_view.showGridLines = False
wb.active = 0
wb.save(OUT); print("ok", OUT, len(datos), "meses de IPC")
